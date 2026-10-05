"""The structured `bli2prism` sheet that sits beside the lab's own `Setup` sheet.

The Setup sheet (plate map, dilution calculator, protein prep table) is left exactly as the
lab keeps it. This sheet adds what a program needs and cannot safely guess:

    [Experiment]     name, results CSV, Prism template, kinetic concentrations per ligand
                     (a manual input), normalization reference analyte, dilution series
    [Plate columns]  one row per plate column: contents (read from Setup), role
                     (Ligand / Reference / Analyte / Buffer / Regeneration), display name,
                     the proteins an entry is made of, and its MW (looked up in Setup)

Yellow cells are inputs. Everything else is a formula, so nothing is typed twice.
"""

from __future__ import annotations

import re

import copy

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

from .xlcalc import Evaluator, fill_cached_values

SHEET = "bli2prism"
SETUP = "Setup"
ROLES = ["Ligand", "Reference", "Analyte", "Buffer", "Regeneration"]
INPUT = PatternFill("solid", fgColor="FFF2CC")
HEAD = PatternFill("solid", fgColor="DCE6F1")
THIN = Side(style="thin", color="B7B7B7")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# fixed keys in the [Experiment] block
KEYS = [
    ("name", "Experiment name", "Free text."),
    ("results_csv", "Results CSV", "File name of the Octet kinetic analysis results table."),
    ("prism_template", "Prism template",
     "Optional. File name of a populated .prism whose graphs are kept (template route). Leave "
     "blank to build without a template."),
    ("prism_output", "Prism output file",
     "Name of the .prism to produce. Upload an EMPTY Prism project with exactly this name into the "
     "run folder and the tool fills it in; the result is written to rebuild/<name>, so your upload "
     "is never overwritten."),
    ("n_conc", "Kinetic concentrations per ligand",
     "REQUIRED, entered by hand: how many dilution steps (A, B, ...) get a kinetics sheet."),
    ("n_analytes", "Number of analytes",
     "REQUIRED, entered by hand (1-7): how many analytes are in the plate map and the raw files. "
     "Must equal the number of plate columns whose Role is Analyte."),
    ("norm_ref", "Normalization reference analyte",
     "Name of the analyte whose MW everything is normalized to (as written in the plate map)."),
    ("dil_start", "Dilution start (M)", "From the Setup sheet."),
    ("dil_factor", "Dilution factor", "From the Setup sheet."),
    ("dil_steps", "Dilution steps", "Concentrations in the series (A..H = 8)."),
]
PLATE_HEAD = ["Plate column", "Contents (from Setup)", "Role", "Display name",
              "Made of (protein 1)", "Made of (protein 2)", "MW (g/mol)", "Check"]
N_PLATE = 12
MAX_ANALYTES = 7
OPTIONAL_KEYS = {"prism_template", "prism_output"}
PROTEIN_NAMES = "ProteinNames"     # defined name for the Proteins tab's name column (a list source Excel keeps)
VOL_PER_ROW_CELL, EXCESS_CELL = "R10", "R11"   # Setup inputs: uL of ligand per sensor row (200), excess factor (1.2)
IMMOB_NAME_CELL, IMMOB_CELL = "Q9", "R9"      # Setup: immobilized ligand concentration (M), an input
PROTEINS = "Proteins"              # tab holding the protein table (A = name, C = MW)
PROT_FIRST, PROT_LAST = 5, 40      # protein table rows (names in A, MWs in C)
PROT_SHIFT = 19                    # legacy Setup layout had the table at T..AG; moved A..N
ROLE_COLORS = {"Ligand": "C6E0B4", "Reference": "D9D2E9", "Analyte": "F8CBAD",
               "Buffer": "EDEDED", "Regeneration": "F4B6C2"}



def q(name):
    """A sheet name as it must be written in a formula ('Plate Map' needs quotes, Setup does not)."""
    return name if re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", name) else "'" + name.replace("'", "''") + "'"


def detect_names(wb):
    """(inputs sheet, plate-map sheet, proteins sheet) titles of a workbook whose tabs may have been renamed,
    found by content: '[Experiment]' in A4, '[Plate map]' in B2, a 'Protein' header in A4. Missing ones fall back to
    the standard names."""
    inputs = plate = prot = None
    for ws in wb.worksheets:
        if ws.cell(4, 1).value == "[Experiment]" and inputs is None:
            inputs = ws.title
        elif ws.cell(2, 2).value == "[Plate map]" and plate is None:
            plate = ws.title
        elif ws.cell(4, 1).value == "Protein" and prot is None:
            prot = ws.title
    return inputs or SHEET, plate or SETUP, prot or PROTEINS


def with_names(fn):
    """Run a writer under the workbook's own tab names (set as the module-level SHEET / SETUP / PROTEINS for the
    duration of the call), so a workbook whose tabs were renamed is handled like any other."""
    import functools

    @functools.wraps(fn)
    def run(wb, *a, **kw):
        global SHEET, SETUP, PROTEINS
        saved = (SHEET, SETUP, PROTEINS)
        SHEET, SETUP, PROTEINS = detect_names(wb)
        try:
            return fn(wb, *a, **kw)
        finally:
            SHEET, SETUP, PROTEINS = saved
    return run


PLATE_FIELDS = [("col", "Plate column", 1), ("contents", "Contents (from Setup)", 2), ("role", "Role", 3),
                ("display", "Display name", 4), ("part1", "Made of (protein 1)", 5),
                ("part2", "Made of (protein 2)", 6), ("mw", "MW (g/mol)", 7), ("check", "Check", 8)]


def plate_columns(ws, header_row):
    """{field: column number} of the plate table, by header text so that inserted or reordered
    columns do not matter; a header that is missing falls back to the standard position."""
    heads = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(header_row, c).value
        if isinstance(v, str):
            heads.setdefault(v.strip(), c)
    return {key: heads.get(label, default) for key, label, default in PLATE_FIELDS}


def plate_rows(ws, title_row):
    """Row numbers of the plate table's body: from the row after the header until the first
    blank 'Plate column' cell or the next [section]. Any number of rows is fine (e.g. several
    ligands stacked on one plate column)."""
    cols = plate_columns(ws, title_row + 1)
    rows, r = [], title_row + 2
    while r <= ws.max_row:
        v = ws.cell(r, cols["col"]).value
        if v in (None, "") or str(v).startswith("["):
            break
        rows.append(r)
        r += 1
    return rows, cols


def _style_input(c):
    c.fill, c.border = INPUT, BOX


SCAN_ROWS = 96                     # 12 plate columns x 8 plate rows
AUTO_SPARE = 6                     # blank rows kept under the listed entries, for plate-map additions
PLATE_ROWS = "ABCDEFGH"


def scan_plate(ws):
    """Entries of the Setup plate map in the order the bli2prism sheet's formulas list them: column by
    column, top to bottom, a name repeated in consecutive cells counted once. -> [(col, first_row,
    last_row, text)], columns 1..12, rows 1..8; an entry lasts until the next entry in its column."""
    out = []
    for k in range(1, N_PLATE + 1):
        texts = []
        for i in range(1, 9):
            v = ws.cell(4 + i, 1 + k).value
            texts.append(v if isinstance(v, str) and v != "" else "" if v is None else str(v))
        starts = [i for i in range(1, 9) if texts[i - 1] and (i == 1 or texts[i - 1] != texts[i - 2])]
        for j, start in enumerate(starts):
            end = starts[j + 1] - 1 if j + 1 < len(starts) else 8
            out.append((k, start, end, texts[start - 1]))
    return out


def _write_scan(ws, head, n_rows, p0):
    """Helper block (columns L..Q) that reads Setup!B5:M12 cell by cell. Returns its ranges."""
    cols = {"k": "L", "row": "M", "name": "N", "start": "O", "idx": "P", "next": "Q"}
    ws[f"L{head - 1}"] = "[Plate map scan]  helper formulas, leave alone"
    ws[f"L{head - 1}"].font = Font(bold=True, color="7F7F7F")
    for col, h in zip("LMNOPQ", ("Plate column", "Row", "Name", "Starts entry", "Entry #", "Next entry at row")):
        c = ws[f"{col}{head}"]
        c.value = h
        c.font, c.fill, c.border = Font(bold=True, color="595959"), HEAD, BOX
        c.alignment = Alignment(wrap_text=True, vertical="center")
    g0 = head + 1
    last = g0 + SCAN_ROWS - 1
    r = g0
    for k in range(1, N_PLATE + 1):
        for i in range(1, 9):
            src = f"{q(SETUP)}!{L(1 + k)}{4 + i}"
            ws[f"L{r}"], ws[f"M{r}"] = k, i
            ws[f"N{r}"] = f'=IF({src}="","",{src})'
            if r == g0:
                ws[f"O{r}"] = f'=IF(N{r}="",0,1)'
                ws[f"P{r}"] = f"=O{r}"
            else:
                ws[f"O{r}"] = f'=IF(N{r}="",0,IF(AND(L{r}=L{r - 1},N{r}=N{r - 1}),0,1))'
                ws[f"P{r}"] = f"=P{r - 1}+O{r}"
            ws[f"Q{r}"] = (9 if r == last else f"=IF(L{r + 1}<>L{r},9,IF(O{r + 1}=1,M{r + 1},Q{r + 1}))")
            for col in "LMNOPQ":
                ws[f"{col}{r}"].font = Font(color="7F7F7F")
            r += 1
    for col, w in zip("LMNOPQ", (12, 6, 34, 12, 10, 14)):
        ws.column_dimensions[col].width = w
    rng = lambda col: f"${col}${g0}:${col}${last}"
    return {c: rng(L_) for c, L_ in (("k", "L"), ("row", "M"), ("name", "N"), ("idx", "P"), ("next", "Q"))}


SENSOR_ROWS_HEAD = "Sensor rows (count)"


def _auto_count_formula(scan, j):
    """Number of plate rows (sensors) the j-th listed entry covers: next entry's start row minus its own."""
    m = f"MATCH({j},{scan['idx']},0)"
    return f'=IFERROR(INDEX({scan["next"]},{m})-INDEX({scan["row"]},{m}),"")'


def _auto_row_formulas(scan, j):
    """(plate column, contents, rows) formulas for the j-th listed plate-map entry."""
    m = f"MATCH({j},{scan['idx']},0)"
    first, nxt = f"INDEX({scan['row']},{m})", f"INDEX({scan['next']},{m})-1"
    letter = lambda x: f'MID("{PLATE_ROWS}",{x},1)'
    return (f'=IFERROR(INDEX({scan["k"]},{m}),"")',
            f'=IFERROR(INDEX({scan["name"]},{m}),"")',
            f'=IFERROR(IF({nxt}={first},{letter(first)},{letter(first)}&"-"&{letter(nxt)}),"")')


def add_sheet(wb, spec):
    """Add the bli2prism sheet. spec: dict with the input values (see draft.infer)."""
    if SHEET in wb.sheetnames:
        del wb[SHEET]
    ws = wb.create_sheet(SHEET, 0)
    ws["A1"] = "bli2prism analysis setup"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = ("Yellow cells are inputs. Contents and MWs are read from the Setup sheet, so "
                "edit the plate map and protein table there.")
    ws["A2"].font = Font(italic=True, color="595959")

    ws["A4"] = "[Experiment]"
    ws["A4"].font = Font(bold=True, color="1F3864")
    row_of = {}
    for i, (key, label, note) in enumerate(KEYS):
        r = 5 + i
        row_of[key] = r
        ws[f"A{r}"], ws[f"C{r}"] = label, note
        ws[f"A{r}"].font = Font(bold=True)
        ws[f"C{r}"].font = Font(italic=True, color="595959")
    v = spec
    for key in ("name", "results_csv", "prism_template", "prism_output", "n_conc", "n_analytes", "norm_ref"):
        if key not in row_of:
            continue
        c = ws[f"B{row_of[key]}"]
        c.value = v.get(key)
        _style_input(c)
    if v.get("analyte_hint") and "n_analytes" in row_of:
        ws[f"C{row_of['n_analytes']}"] = (ws[f"C{row_of['n_analytes']}"].value
                                          + f"  (Detected {v['analyte_hint']} Analyte column(s) in the plate map.)")
    for key, hi in (("n_conc", 8), ("n_analytes", MAX_ANALYTES)):
        if key not in row_of:
            continue
        dvw = DataValidation(type="whole", operator="between", formula1="1", formula2=str(hi),
                             allow_blank=True, showErrorMessage=True,
                             errorTitle="Whole number", error=f"Enter a whole number from 1 to {hi}.")
        ws.add_data_validation(dvw)
        dvw.add(f"B{row_of[key]}")
    ws[f"B{row_of['dil_start']}"] = f"={q(SETUP)}!N5"
    ws[f"B{row_of['dil_factor']}"] = f"={q(SETUP)}!R5"
    ws[f"B{row_of['dil_steps']}"] = v.get("dil_steps", 8)
    _style_input(ws[f"B{row_of['dil_steps']}"])
    ws[f"B{row_of['dil_start']}"].number_format = "0.00E+00"

    title = 5 + len(KEYS) + 1
    ws[f"A{title}"] = "[Plate columns]"
    ws[f"A{title}"].font = Font(bold=True, color="1F3864")
    head = title + 1
    for i, h in enumerate(PLATE_HEAD, 1):
        c = ws.cell(head, i, h)
        c.font, c.fill, c.border = Font(bold=True), HEAD, BOX
        c.alignment = Alignment(wrap_text=True, vertical="center")
    p0 = head + 1
    names = f"{q(PROTEINS)}!$A${PROT_FIRST}:$A${PROT_LAST}"
    mws = f"{q(PROTEINS)}!$C${PROT_FIRST}:$C${PROT_LAST}"
    plate = v.get("plate", {})
    auto = bool(v.get("auto"))
    if auto:
        entries = [plate[k] for k in sorted(plate)]
        entries += [{} for _ in range(max(AUTO_SPARE, N_PLATE - len(entries)))]
        entries = entries[:SCAN_ROWS]
        scan = _write_scan(ws, head, len(entries), p0)
        for col, h in (("I", "Plate rows"), ("J", "Label (shown on Setup)"), ("K", SENSOR_ROWS_HEAD)):
            c = ws[f"{col}{head}"]
            c.value = h
            c.font, c.fill, c.border = Font(bold=True), HEAD, BOX
            c.alignment = Alignment(wrap_text=True, vertical="center")
    else:
        entries = [plate[k] for k in sorted(plate)] if plate else [{} for _ in range(N_PLATE)]
    n_rows = len(entries)
    for idx, entry in enumerate(entries):
        k = entry.get("col", idx + 1)
        r = p0 + idx
        if auto:
            ws[f"A{r}"], ws[f"B{r}"], ws[f"I{r}"] = _auto_row_formulas(scan, idx + 1)
            ws[f"J{r}"] = (f'=IF(C{r}="","",IF(OR(C{r}="Ligand",C{r}="Analyte"),'
                           f'C{r}&" "&COUNTIFS($C${p0}:C{r},C{r}),C{r}))')
            ws[f"K{r}"] = _auto_count_formula(scan, idx + 1)
            for col in "IJK":
                ws[f"{col}{r}"].border = BOX
        else:
            ws[f"A{r}"] = k
            ws[f"B{r}"] = entry.get("contents_formula") or f'=IF({q(SETUP)}!{L(1 + k)}5="","",{q(SETUP)}!{L(1 + k)}5)'
        for col, key in (("C", "role"), ("D", "display"), ("E", "part1"), ("F", "part2")):
            c = ws[f"{col}{r}"]
            c.value = entry.get(key)
            _style_input(c)
        ws[f"G{r}"] = (f'=IF(OR(C{r}="Ligand",C{r}="Reference",C{r}="Analyte"),'
                       f'IF(E{r}="",0,SUMIF({names},E{r},{mws}))'
                       f'+IF(F{r}="",0,SUMIF({names},F{r},{mws})),"")')
        ws[f"G{r}"].number_format = "#,##0.00"
        ws[f"H{r}"] = entry.get("check", "")
        for col in "ABGH":
            ws[f"{col}{r}"].border = BOX
        ws[f"H{r}"].font = Font(italic=True, color="C00000")
    add_role_colors(ws, f"C{p0}:C{p0 + n_rows - 1}", f"$C{p0}", exact=True)
    dv = DataValidation(type="list", formula1='"' + ",".join(ROLES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"C{p0}:C{p0 + n_rows - 1}")
    dv2 = DataValidation(type="list", formula1=f"=$B${p0}:$B${p0 + n_rows - 1}", allow_blank=True)
    ws.add_data_validation(dv2)
    dv2.add(f"B{row_of['norm_ref']}")
    # protein pickers: names come from the Proteins tab, so the MW lookup in column G cannot miss on a typo
    _define_protein_names(wb)
    dv_prot = DataValidation(type="list", formula1=f"={PROTEIN_NAMES}", allow_blank=True, showErrorMessage=True,
                             errorTitle="Not in the Proteins tab",
                             error="Pick a protein from the list (add it to the Proteins tab first if it is missing).")
    ws.add_data_validation(dv_prot)
    dv_prot.add(f"E{p0}:F{p0 + n_rows - 1}")
    # display names: suggestions are the plate-map names above, but any text is accepted (e.g. with Greek letters)
    dv_disp = DataValidation(type="list", formula1=f"=$B${p0}:$B${p0 + n_rows - 1}", allow_blank=True,
                             showErrorMessage=False)
    ws.add_data_validation(dv_disp)
    dv_disp.add(f"D{p0}:D{p0 + n_rows - 1}")
    if auto:
        ws[f"A{title}"] = "[Plate columns]"
        ws[f"C{title}"] = ("Columns A, B, I, J are formulas that list every entry of the Setup plate map: edit the plate "
                           "map, not these. Role, Display name and Made of are yours to fill in; if you reorder the plate "
                           "map, re-check them, they do not move with the names.")
        ws[f"C{title}"].font = Font(italic=True, color="595959")
    for col, w in zip("ABCDEFGHIJK", (34, 36, 14, 30, 28, 28, 14, 48, 12, 22, 14)):
        ws.column_dimensions[col].width = w
    ws.sheet_view.showGridLines = False
    return ws


# ------------------------------------------------------------ Setup sheet styling
SECTION = Font(bold=True, color="1F3864")
NOTE = Font(italic=True, color="595959", size=9)
BODY = "Calibri"
PLAIN = PatternFill(fill_type=None)


def plate_row0():
    """First plate-column row on the bli2prism sheet (where each column's Role lives)."""
    return 5 + len(KEYS) + 3


def add_role_colors(ws, ref, anchor, exact=False):
    """Conditional fills per role. anchor = the top-left cell of the rule, role text in it."""
    from openpyxl.formatting.rule import FormulaRule
    for role, color in ROLE_COLORS.items():
        test = f'{anchor}="{role}"' if exact or role not in ("Ligand", "Analyte") \
            else f'LEFT({anchor},{len(role)})="{role}"'
        ws.conditional_formatting.add(ref, FormulaRule(
            formula=[test], fill=PatternFill(start_color=color, end_color=color, fill_type="solid")))


def _box_range(ws, ref, fill=None, bold=None, wrap=False, center=False):
    """Style every cell of a range (merged ranges need each cell for the borders to show)."""
    rows = ws[ref] if ":" in ref else ((ws[ref],),)
    for row in rows:
        for c in row:
            c.border = BOX
            if fill is not None:
                c.fill = fill
            if bold is not None:
                c.font = Font(name=BODY, size=11, bold=bold)
            if wrap or center:
                c.alignment = Alignment(horizontal="center" if center else None,
                                        vertical="center", wrap_text=wrap)


def _clear(c):
    c.value = None
    c.style = "Normal"


def split_proteins(wb):
    """Move the protein / bench-prep block from the Setup sheet (T4:AG..) to its own tab
    (A4:N..), translating its formulas and repointing every formula elsewhere in the workbook
    that looked values up in the old place. Returns True if anything moved."""
    from openpyxl.formula.translate import Translator
    from openpyxl.utils import get_column_letter, range_boundaries
    ws = wb[SETUP]
    if PROTEINS in wb.sheetnames or ws["T4"].value != "Protein":
        return False
    new = wb.create_sheet(PROTEINS)
    last = max(r for r in range(4, 80) if ws.cell(r, 20).value)
    c0, c1 = 20, 33
    for r in range(4, last + 1):
        for col in range(c0, c1 + 1):
            src, dst = ws.cell(r, col), new.cell(r, col - PROT_SHIFT)
            v = src.value
            if isinstance(v, str) and v.startswith("="):
                v = Translator(v, origin=src.coordinate).translate_formula(dst.coordinate)
                # absolute references to the dilution block now live on another sheet
                v = re.sub(r"(?<![!A-Za-z])(\$N\$5|\$R\$8)", lambda m: f"{q(SETUP)}!{m.group(1)}", v)
            dst.value = v
            dst.number_format = src.number_format
    for rng in list(ws.merged_cells.ranges):
        mc0, mr0, mc1, mr1 = range_boundaries(str(rng))
        if mc0 >= c0 and mr0 >= 4:
            ws.unmerge_cells(str(rng))
            new.merge_cells(start_row=mr0, start_column=mc0 - PROT_SHIFT,
                            end_row=mr1, end_column=mc1 - PROT_SHIFT)
    for r in range(2, last + 1):
        for col in range(c0, c1 + 1):
            _clear(ws.cell(r, col))
    for col in range(c0, c1 + 1):
        ws.column_dimensions.pop(get_column_letter(col), None)
    # repoint lookups elsewhere (e.g. the bli2prism sheet's MW formulas)
    pat = re.compile(rf"{q(SETUP)}!(\$?)([A-Z]{{1,2}})(\$?)(\d+)(?::(\$?)([A-Z]{{1,2}})(\$?)(\d+))?")
    from openpyxl.utils import column_index_from_string as ci

    def fix(m):
        c_a = ci(m.group(2))
        if c_a < c0 or c_a > c1:
            return m.group(0)
        out = f"{q(PROTEINS)}!{m.group(1)}{get_column_letter(c_a - PROT_SHIFT)}{m.group(3)}{m.group(4)}"
        if m.group(6):
            out += f":{m.group(5)}{get_column_letter(ci(m.group(6)) - PROT_SHIFT)}{m.group(7)}{m.group(8)}"
        return out
    for sh in wb.worksheets:
        if sh.title == PROTEINS:
            continue
        for row in sh.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and f"{q(SETUP)}!" in c.value:
                    c.value = pat.sub(fix, c.value)
    return True


def style_proteins(ws):
    """Style the Proteins tab like the others: blue headers, yellow inputs, white formulas."""
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                c.font = Font(name=BODY, size=11, bold=c.font.bold, italic=c.font.italic)
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Proteins and bench prep"
    ws["A1"].font = Font(name=BODY, size=14, bold=True)
    ws["C1"] = ("Yellow cells are inputs. White cells are formulas: do not overwrite them. The "
                "bli2prism sheet reads protein names (column A) and MWs (column C) in rows 5-40.")
    ws["C1"].font = Font(name=BODY, italic=True, color="595959")
    ws["A2"] = ("Inputs: name, A280, MW, extinction coefficient, and the desired concentration / total "
                "volume where they are typed. To add a protein, copy the last row down.")
    ws["A2"].font = NOTE
    _box_range(ws, "A4:N4", HEAD, bold=True, center=True, wrap=True)
    last = max((r for r in range(5, 80) if ws.cell(r, 1).value), default=11)
    for r in range(5, last + 1):
        for col in range(1, 15):
            c = ws.cell(r, col)
            is_formula = isinstance(c.value, str) and c.value.startswith("=")
            is_input = col <= 4 or (col in (9, 10) and not is_formula and c.value is not None)
            c.border = BOX
            c.fill = INPUT if is_input else PLAIN
            if c.value is not None or is_input:
                c.font = Font(name=BODY, size=11)
    ws.row_dimensions[4].height = 48
    widths = [26, 11, 13, 18, 18, 18, 20, 16, 20, 18, 16, 16, 18, 20]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[L(i)].width = w
    ws.freeze_panes = "B5"


def style_setup(ws, auto_rows=None):
    """Give the Setup sheet the bli2prism sheet's look and label each plate column's role.

    Formatting plus annotation: values and formulas already on the sheet are kept (the one
    exception is the merged 'Immobilized proteins' header over the ligand columns, whose
    wording moves into a note so row 3 can carry per-column roles). With auto_rows (the number of
    plate-table rows on an auto-listing bli2prism sheet) row 3 looks each column's role up by plate
    column instead of by position; a stacked column shows its first entry. Yellow = type here;
    white = formula or label; blue = header; role chips are coloured by role.
    """
    for row in ws.iter_rows():                       # drop legacy fills, unify the font
        for c in row:
            if c.fill.fill_type:
                c.fill = PLAIN
            if c.value is not None:
                c.font = Font(name=BODY, size=11, bold=c.font.bold, italic=c.font.italic)
    ws.sheet_view.showGridLines = False
    from openpyxl.formatting.formatting import ConditionalFormattingList
    ws.conditional_formatting = ConditionalFormattingList()     # re-styling must not stack rules

    ws["A1"] = "Setup"
    ws["A1"].font = Font(name=BODY, size=14, bold=True)
    ws["C1"] = ("Yellow cells are inputs. White cells are formulas or labels: do not overwrite them. "
                "The bli2prism sheet reads this sheet, so edit values but do not move cells.")
    ws["C1"].font = Font(name=BODY, italic=True, color="595959")
    for ref, text in (("B2", "[Plate map]"), ("N2", "[Dilution series]")):
        ws[ref] = text
        ws[ref].font = Font(name=BODY, size=11, bold=True, color="1F3864")
    ws["N3"] = "Starting concentration, dilution factor and final volume are inputs."
    ws["N3"].font = NOTE

    # --- plate map: per-column role row (row 3) mirrored from the bli2prism sheet
    for rng in list(ws.merged_cells.ranges):
        if str(rng) == "C3:E3":
            ws.unmerge_cells("C3:E3")
    old = ws["C3"].value
    if isinstance(old, str) and old.startswith("Immobilized"):
        ws["C3"].value = None                          # wording kept in the note below
    ws["D2"] = ("Immobilized proteins (100 nM) are the Ligand and Reference columns. "
                "Roles are set on the bli2prism tab (Role column).")
    ws["D2"].font = NOTE
    if ws["G3"].value and str(ws["G3"].value).startswith("The tool reads plate row A"):
        ws["G3"].value = None
    p0 = plate_row0()
    ws["A3"], ws["A4"] = "Role", "Column"
    _box_range(ws, "A3:A4", HEAD, bold=True, center=True)
    for k in range(1, 13):
        if auto_rows:
            last = p0 + auto_rows - 1
            pick = (f"INDEX({q(SHEET)}!$J${p0}:$J${last},MATCH({k},{q(SHEET)}!$A${p0}:$A${last},0))")
            more = f'COUNTIFS({q(SHEET)}!$A${p0}:$A${last},{k})-1'
            ws.cell(3, 1 + k).value = (f'=IFERROR(IF({pick}="","(set role)",{pick}'
                                       f'&IF({more}>0," +"&{more}&" more","")),"")')
            continue
        r = p0 + k - 1
        ref = f"{q(SHEET)}!C{r}"
        ws.cell(3, 1 + k).value = (
            f'=IF({ref}="","(set role)",IF(OR({ref}="Ligand",{ref}="Analyte"),'
            f'{ref}&" "&COUNTIFS({q(SHEET)}!$C${p0}:C{r},{ref}),{ref}))')
    _box_range(ws, "B3:M3", None, bold=True, center=True, wrap=True)
    ws.row_dimensions[3].height = 32
    add_role_colors(ws, "B3:M4", "B$3")
    from openpyxl.formatting.rule import FormulaRule
    ws.conditional_formatting.add("B3:M3", FormulaRule(
        formula=['B$3="(set role)"'], font=Font(italic=True, color="9A9A9A")))

    _box_range(ws, "B4:M4", HEAD, bold=True, center=True)         # column numbers
    _box_range(ws, "A5:A12", HEAD, bold=True, center=True)        # row letters
    for col in range(2, 14):                                      # contents: inputs
        top = ws.cell(5, col)
        merged = next((r for r in ws.merged_cells.ranges if top.coordinate in r), None)
        ref = str(merged) if merged else f"{L(col)}5:{L(col)}12"
        _box_range(ws, ref, INPUT, bold=False, wrap=True, center=True)
    ws.row_dimensions[4].height = 22
    ws["B13"] = "The tool reads plate row A (row 5); the other rows are for your records."
    ws["B14"] = ("Ligand N = raw files N.xls (1.xls, 2.xls, ...); Analyte N = the Nth trace in each "
                 "raw file; Reference = the sensor that was subtracted.")
    for ref in ("B13", "B14"):
        ws[ref].font = NOTE

    # --- dilution series
    _box_range(ws, "N4:P4", HEAD, bold=True, center=True, wrap=True)
    _box_range(ws, "N5:P12", None, bold=False)
    _box_range(ws, "N5", INPUT)                                    # starting concentration
    _box_range(ws, "Q4:R4", HEAD, bold=True, center=True)
    ws["Q4"] = "Dilution settings"
    if "Q4:R4" not in {str(r) for r in ws.merged_cells.ranges}:
        ws.merge_cells("Q4:R4")
    _box_range(ws, "Q5:Q8", None, bold=True)
    _box_range(ws, "R5:R8", None, bold=False)
    for ref in ("R5", "R7"):                                       # factor, final volume
        _box_range(ws, ref, INPUT, bold=False)
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 9)


def explain_ligand_mismatch(setup, loading_ids):
    """A message for when the sheet's ligands differ from the results CSV's, naming the plate-table
    rows involved and the likely cause (a Contents formula pointing at the wrong Setup cell)."""
    csv_ids = set(loading_ids)
    only_csv = sorted(csv_ids - set(setup["ligands"]))
    only_sheet = [x for x in setup["ligands"] if x not in csv_ids]
    lines = ["the bli2prism sheet and the results CSV disagree about the ligands:"]
    if only_csv:
        lines.append(f"  in the results CSV but not on the sheet: {only_csv}")
    if only_sheet:
        lines.append(f"  on the sheet but not in the CSV: {[x or '(empty)' for x in only_sheet]}")
    for e in setup["plate"]:
        if e["role"] != "Ligand":
            continue
        made = e["parts"][0] if e["parts"] else None
        unknown = e["contents"] not in csv_ids
        disagrees = bool(made) and made in csv_ids and e["contents"] != made
        if not (unknown or disagrees):
            continue
        what = f"reads {e['contents']!r}" if e["contents"] else "is empty"
        via = f" (formula {e['contents_formula']})" if e["contents_formula"] else ""
        hint = ""
        if disagrees:
            hint = (f"; 'Made of (protein 1)' says {made!r}, which is in the CSV, so this Contents formula "
                    "probably points at the wrong Setup plate-map cell")
        lines.append(f"  plate table row {e['row']} (Role = Ligand): Contents {what}{via}{hint}")
    lines.append("Fix the Contents cells, or run `bli2prism fix-contents <setup file>` to point them at the "
                 "Setup plate-map cells holding the names in 'Made of (protein 1)'.")
    return "\n".join(lines)


def auto_row_count(ws):
    """Number of rows in the plate table of a sheet that add_sheet just wrote."""
    labels = {r[0].value: r[0].row for r in ws.iter_rows(min_col=1, max_col=1) if isinstance(r[0].value, str)}
    body, _ = plate_rows(ws, labels["[Plate columns]"])
    return len(body)


def is_auto(ws, cols, body):
    """True when the plate table is the auto-listing kind (its Plate column cells are formulas)."""
    v = ws.cell(body[0], cols["col"]).value
    return isinstance(v, str) and v.startswith("=")


@with_names
def fix_contents(wb):
    """Re-point wrong Contents formulas of Ligand / Reference rows at the Setup plate-map cell that
    holds the name written in 'Made of (protein 1)'. Only rows whose Contents currently differs from
    that name, and only when the name occurs in exactly one plate-map cell, are changed.
    Returns (changes, skipped): lists of (row, old formula, new formula) and (row, reason)."""
    ws, setup = wb[SHEET], wb[SETUP]
    labels = {r[0].value: r[0].row for r in ws.iter_rows(min_col=1, max_col=1) if isinstance(r[0].value, str)}
    body, cols = plate_rows(ws, labels["[Plate columns]"])
    if body and is_auto(ws, cols, body):
        raise ValueError("this bli2prism sheet lists the plate map by formula, so its Contents cells cannot "
                         "point at the wrong Setup cell: nothing to fix")
    grid = {}
    for r in range(5, 13):
        for c in range(2, 14):
            v = setup.cell(r, c).value
            if isinstance(v, str) and v.strip() and not v.startswith("="):
                grid.setdefault(v.strip(), []).append((r, c))
    ev = Evaluator(wb)
    changes, skipped = [], []
    for r in body:
        role = ws.cell(r, cols["role"]).value
        if role not in ("Ligand", "Reference"):
            continue
        want = ws.cell(r, cols["part1"]).value
        if not want:
            skipped.append((r, f"Role is {role} but 'Made of (protein 1)' is empty"))
            continue
        try:
            current = ev.value(f"{L(cols['contents'])}{r}", SHEET)
        except Exception:
            current = None
        if current == want:
            continue
        spots = grid.get(want.strip(), [])
        if len({c for _, c in spots}) != 1:          # several columns: cannot tell which one is meant
            skipped.append((r, f"{want!r} " + ("is not in the Setup plate map" if not spots else
                                                f"appears in {len({c for _, c in spots})} plate-map columns, cannot choose")))
            continue
        spots = [min(spots)]                          # repeated down one column: the top of the block
        ref = f"{q(SETUP)}!{L(spots[0][1])}{spots[0][0]}"
        old = ws.cell(r, cols["contents"]).value
        new = f'=IF({ref}="","",{ref})'
        ws.cell(r, cols["contents"]).value = new
        changes.append((r, old, new))
    return changes, skipped


@with_names
def upgrade_sheet(wb):
    """Add inputs that a bli2prism sheet made by an older version lacks ('Number of analytes',
    'Prism output file'). The sheet is laid
    out again (the plate table moves down a row) with every input kept, and the Setup tab's
    role row is re-pointed. Returns True if anything changed."""
    if SHEET not in wb.sheetnames:
        raise ValueError("no bli2prism sheet in this workbook")
    ws = wb[SHEET]
    labels = {r[0].value: r[0].row for r in ws.iter_rows(min_col=1, max_col=1) if isinstance(r[0].value, str)}
    if all(label in labels for _, label, _ in KEYS):
        return False
    spec = {}
    for key, label, _ in KEYS:
        if label in labels:
            v = ws[f"B{labels[label]}"].value
            if not (isinstance(v, str) and v.startswith("=")):
                spec[key] = v
    body, cols = plate_rows(ws, labels["[Plate columns]"])
    if body and is_auto(ws, cols, body):
        raise ValueError("this bli2prism sheet lists the plate map by formula and cannot be upgraded in place: "
                         "draft a new one with `bli2prism draft`")
    plate = {}
    for k, r in enumerate(body, 1):
        cf = ws.cell(r, cols["contents"]).value
        plate[k] = {"col": ws.cell(r, cols["col"]).value,
                    "contents_formula": cf if isinstance(cf, str) and cf.startswith("=") else None,
                    "role": ws.cell(r, cols["role"]).value, "display": ws.cell(r, cols["display"]).value,
                    "part1": ws.cell(r, cols["part1"]).value, "part2": ws.cell(r, cols["part2"]).value,
                    "check": ws.cell(r, cols["check"]).value or ""}
    spec["plate"] = plate
    spec["analyte_hint"] = sum(1 for e in plate.values() if e["role"] == "Analyte")
    add_sheet(wb, spec)
    if SETUP in wb.sheetnames:
        style_setup(wb[SETUP])
    return True


def _define_protein_names(wb):
    """Workbook-level name for Proteins!A5:A40. Dropdowns that point at another sheet directly are rewritten by
    Excel into a form openpyxl drops on the next save; a named range stays an ordinary list source."""
    from openpyxl.workbook.defined_name import DefinedName
    ref = f"{q(PROTEINS)}!$A${PROT_FIRST}:$A${PROT_LAST}"
    if PROTEIN_NAMES in wb.defined_names:
        del wb.defined_names[PROTEIN_NAMES]
    wb.defined_names[PROTEIN_NAMES] = DefinedName(PROTEIN_NAMES, attr_text=ref)


@with_names
def link_proteins(wb, immobilized_M=1e-7):
    """Make the Proteins tab follow the plate map's roles, and give Setup an immobilized-ligand concentration input.

    Setup!R9 (input, M) holds the concentration the ligand and reference sensors were loaded at. On the Proteins tab a
    new column N reads each protein's role from the bli2prism sheet (a protein named under 'Made of' on a Ligand,
    Reference or Analyte row), and 'Desired Concentration (mol/L)' becomes a formula: immobilized concentration for
    Ligand / Reference proteins, the Setup starting concentration for Analytes, 0 (shown blank) for a protein that is
    unused. A protein that is both a ligand and an analyte is prepared for the analyte concentration (it needs two
    preps). Moles Needed is rewritten per row (one row of the lab's sheet pointed at the row above). The 'Made of'
    dropdowns are re-asserted. Returns a list of what changed.
    """
    missing = [n for n in (SHEET, SETUP, PROTEINS) if n not in wb.sheetnames]
    if missing:
        raise ValueError("could not find the inputs sheet ('[Experiment]' in A4), the plate-map sheet ('[Plate map]' in "
                         f"B2) and the Proteins sheet ('Protein' in A4) in this workbook (tabs: {wb.sheetnames})")
    bl, st, pr = wb[SHEET], wb[SETUP], wb[PROTEINS]
    labels = {r[0].value: r[0].row for r in bl.iter_rows(min_col=1, max_col=1) if isinstance(r[0].value, str)}
    body, cols = plate_rows(bl, labels["[Plate columns]"])
    p0 = body[0]
    last_pl = p0 + max(len(body), SCAN_ROWS) - 1
    rng = lambda key: f"{q(SHEET)}!${L(cols[key])}${p0}:${L(cols[key])}${last_pl}"
    done = []

    # Setup: the input
    if st[IMMOB_CELL].value in (None, ""):
        st[IMMOB_CELL] = immobilized_M
    st[IMMOB_NAME_CELL] = "Immobilized ligand concentration (M)"
    st[IMMOB_NAME_CELL]._style = copy.copy(st["Q8"]._style)
    st[IMMOB_CELL]._style = copy.copy(st["R7"]._style)       # an input cell (yellow)
    st[IMMOB_CELL].number_format = "0.00E+00"
    st.column_dimensions["Q"].width = max(st.column_dimensions["Q"].width or 0, 34)
    d2 = st["D2"].value
    if isinstance(d2, str) and "(100 nM)" in d2:
        st["D2"] = d2.replace("(100 nM)", f"(concentration: {IMMOB_CELL})")
    for name_cell, cell, text, default in (("Q10", VOL_PER_ROW_CELL, "Ligand volume per sensor row (uL)", 200),
                                           ("Q11", EXCESS_CELL, "Volume excess factor", 1.2)):
        if st[cell].value in (None, ""):
            st[cell] = default
        st[name_cell] = text
        st[name_cell]._style = copy.copy(st["Q8"]._style)
        st[cell]._style = copy.copy(st["R7"]._style)
    st[IMMOB_CELL]._style = copy.copy(st["R7"]._style)
    done.append(f"Setup!{IMMOB_CELL}: immobilized ligand concentration input ({st[IMMOB_CELL].value:g} M); "
                f"{VOL_PER_ROW_CELL}: ligand volume per sensor row; {EXCESS_CELL}: excess factor")

    # Proteins: role column N and desired concentration H
    first, last = PROT_FIRST, max((r for r in range(PROT_FIRST, PROT_LAST + 1)
                                   if isinstance(pr.cell(r, 5).value, str) and pr.cell(r, 5).value.startswith("=")),
                                  default=PROT_FIRST)
    for mr in list(pr.merged_cells.ranges):          # a merged Total Volume cell cannot hold a formula per protein
        if mr.min_col <= 9 <= mr.max_col and mr.max_row >= first:
            pr.unmerge_cells(str(mr))
    pr["N4"] = "Role (from bli2prism tab)"
    pr["N4"]._style = copy.copy(pr["M4"]._style)
    pr.column_dimensions["N"].width = 26

    head = labels["[Plate columns]"] + 1
    heads = {bl.cell(head, c).value: c for c in range(1, bl.max_column + 1)}
    n_col = heads.get(SENSOR_ROWS_HEAD)
    if not n_col and is_auto(bl, cols, body):
        # an auto sheet drafted before the count column existed: add it in the first empty column after the label
        # column, with formulas over the scan block, all found by header (columns may have been moved or deleted)
        scan_head = {bl.cell(head, c).value: c for c in range(1, bl.max_column + 1)
                     if c > cols["check"] and bl.cell(head, c).value}
        need = {"Row", "Name", "Next entry at row", "Entry #"}
        label_col = heads.get("Label (shown on Setup)")
        if label_col and need <= set(scan_head) and "Plate column" in scan_head:
            n_col = next(c for c in range(label_col + 1, bl.max_column + 2) if bl.cell(head, c).value in (None, ""))
            g0 = head + 1
            rg = lambda c: f"${L(c)}${g0}:${L(c)}${g0 + SCAN_ROWS - 1}"
            scan = {"k": rg(scan_head["Plate column"]), "row": rg(scan_head["Row"]), "name": rg(scan_head["Name"]),
                    "idx": rg(scan_head["Entry #"]), "next": rg(scan_head["Next entry at row"])}
            bl[f"{L(n_col)}{head}"] = SENSOR_ROWS_HEAD
            bl[f"{L(n_col)}{head}"]._style = copy.copy(bl.cell(head, label_col)._style)
            for j, r in enumerate(body, 1):
                bl[f"{L(n_col)}{r}"] = _auto_count_formula(scan, j)
                bl[f"{L(n_col)}{r}"]._style = copy.copy(bl.cell(r, label_col)._style)
            bl.column_dimensions[L(n_col)].width = 14
            done.append(f"bli2prism!{L(n_col)}{body[0]}:{L(n_col)}{body[-1]}: sensor rows (count) per plate-map entry added")
    if n_col:
        n_rng = f"{q(SHEET)}!${L(n_col)}${p0}:${L(n_col)}${last_pl}"

        def sensors(role, r):
            return (f'SUMIFS({n_rng},{rng("role")},"{role}",{rng("part1")},$A{r})'
                    f'+SUMIFS({n_rng},{rng("role")},"{role}",{rng("part2")},$A{r})')
    else:                                           # a static table does not know how many rows: assume the full 8
        def sensors(role, r):
            return f'8*({count(role, r)})'

    def count(role, r):
        return (f'COUNTIFS({rng("role")},"{role}",{rng("part1")},$A{r})'
                f'+COUNTIFS({rng("role")},"{role}",{rng("part2")},$A{r})')
    for r in range(first, last + 1):
        lig, ana, ref = count("Ligand", r), count("Analyte", r), count("Reference", r)
        pr[f"N{r}"] = (f'=IF($A{r}="","",IF(AND({lig}>0,{ana}>0),"Ligand + Analyte",IF({ana}>0,"Analyte",'
                       f'IF({lig}>0,"Ligand",IF({ref}>0,"Reference","")))))')
        pr[f"H{r}"] = (f'=IF(OR(N{r}="Analyte",N{r}="Ligand + Analyte"),{q(SETUP)}!$N$5,'
                       f'IF(OR(N{r}="Ligand",N{r}="Reference"),{q(SETUP)}!${IMMOB_CELL[0]}${IMMOB_CELL[1:]},0))')
        pr[f"G{r}"] = f"=H{r}*I{r}*10^-6"
        vol_lig = f"({sensors('Ligand', r)}+{sensors('Reference', r)})*{q(SETUP)}!${VOL_PER_ROW_CELL[0]}${VOL_PER_ROW_CELL[1:]}"
        pr[f"I{r}"] = (f'=IF(OR(N{r}="Analyte",N{r}="Ligand + Analyte"),{q(SETUP)}!$R$8*{q(SETUP)}!${EXCESS_CELL[0]}${EXCESS_CELL[1:]},'
                       f'IF(OR(N{r}="Ligand",N{r}="Reference"),{vol_lig}*{q(SETUP)}!${EXCESS_CELL[0]}${EXCESS_CELL[1:]},0))')
        pr[f"I{r}"].number_format = "#,##0;-#,##0;"
        for col in "GHIN":
            pr[f"{col}{r}"]._style = copy.copy(pr[f"E{r}"]._style)      # white formula cell
        pr[f"H{r}"].number_format = "0.00E+00;-0.00E+00;"              # unused proteins show blank
    pr["A2"] = ("Inputs: name, A280, MW and extinction coefficient. Role comes from the bli2prism tab; the desired "
                f"concentration and total volume follow it (see README: Ligand / Reference use Setup!{IMMOB_CELL} and the sensor rows, Analyte uses Setup!N5 and R8). To add a protein, copy the last row down.")
    done.append(f"Proteins!N{first}:N{last}: role by formula; H{first}:H{last}: desired concentration by role; "
                f"G{first}:G{last}: moles needed re-pointed at its own row")

    # dropdowns on Made of (protein 1 / 2), by the named range
    _define_protein_names(wb)
    keep = []
    for dv in bl.data_validations.dataValidation:
        ranges = [r for r in dv.sqref.ranges if not (r.min_col in (cols["part1"], cols["part2"]) and r.min_row >= p0)]
        if ranges:
            dv.sqref = type(dv.sqref)(" ".join(str(r) for r in ranges))
            keep.append(dv)
    bl.data_validations.dataValidation = keep
    dvp = DataValidation(type="list", formula1=f"={PROTEIN_NAMES}", allow_blank=True, showErrorMessage=True,
                         errorTitle="Not in the Proteins tab",
                         error="Pick a protein from the list (add it to the Proteins tab first if it is missing).")
    bl.add_data_validation(dvp)
    dvp.add(f"{L(cols['part1'])}{p0}:{L(cols['part2'])}{last_pl}")
    done.append("bli2prism: Made of dropdowns re-asserted (named range ProteinNames)")
    return done


# ------------------------------------------------------------------------- reader
def sheet_title(wb):
    """Title of the structured input sheet: 'bli2prism', or whatever it was renamed to (found by its
    '[Experiment]' block in A4). None if the workbook has no such sheet."""
    if SHEET in wb.sheetnames:
        return SHEET
    for ws in wb.worksheets:
        if ws.cell(4, 1).value == "[Experiment]":
            return ws.title
    return None


def has_sheet(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return sheet_title(wb) is not None
    finally:
        wb.close()


def read_value(path, key):
    """One [Experiment] input by key (e.g. 'results_csv'), without validating the rest."""
    wf = openpyxl.load_workbook(path)
    wv = openpyxl.load_workbook(path, data_only=True)
    title = sheet_title(wf)
    ws, wsv = wf[title], wv[title]
    label = next(l for k, l, _ in KEYS if k == key)
    for row in ws.iter_rows(min_col=1, max_col=1):
        if row[0].value == label:
            addr = f"B{row[0].row}"
            v = wsv[addr].value
            if v is None and isinstance(ws[addr].value, str) and ws[addr].value.startswith("="):
                v = Evaluator(wf).value(addr, title)
            return v
    return None


def read(path):
    """-> dict with ligands, reference, analytes, series_M, n_concentrations, ... (validated)."""
    wf = openpyxl.load_workbook(path)
    wv = openpyxl.load_workbook(path, data_only=True)
    ev = Evaluator(wf)
    title = sheet_title(wf)
    ws, wsv = wf[title], wv[title]

    def val(addr):
        v = wsv[addr].value
        if v is None and isinstance(ws[addr].value, str) and ws[addr].value.startswith("="):
            try:
                v = ev.value(addr, title)
            except Exception:
                v = None
        return v

    labels = {row[0].value: row[0].row for row in ws.iter_rows(min_col=1, max_col=1)
              if isinstance(row[0].value, str)}
    exp = {}
    for key, label, _ in KEYS:
        if key in OPTIONAL_KEYS and label not in labels:
            exp[key] = None                      # a sheet from before this input existed still works
            continue
        if label not in labels:
            raise ValueError(f"{title} sheet: missing '{label}' in the [Experiment] block. This file "
                             "was made by an older version: run `bli2prism upgrade <file>` to add it")
        exp[key] = val(f"B{labels[label]}")
    if "[Plate columns]" not in labels:
        raise ValueError(f"{title} sheet: missing the [Plate columns] block")
    body, cols = plate_rows(ws, labels["[Plate columns]"])
    letter = lambda key: L(cols[key])
    plate = []
    for r in body:
        cf = ws[f"{letter('contents')}{r}"].value
        plate.append({"row": r, "col": val(f"{letter('col')}{r}"),
                      "contents": val(f"{letter('contents')}{r}") or "",
                      "contents_formula": cf if isinstance(cf, str) and cf.startswith("=") else None,
                      "role": val(f"{letter('role')}{r}"), "display": val(f"{letter('display')}{r}"),
                      "parts": [p for p in (val(f"{letter('part1')}{r}"), val(f"{letter('part2')}{r}")) if p],
                      "mw": val(f"{letter('mw')}{r}")})

    problems = []
    ligands = [p["contents"] for p in plate if p["role"] == "Ligand"]
    refs = [p["contents"] for p in plate if p["role"] == "Reference"]
    analytes = []
    for p in plate:
        if p["role"] != "Analyte":
            continue
        if not isinstance(p["mw"], (int, float)) or p["mw"] <= 0:
            problems.append(f"analyte '{p['contents']}' has no MW: set 'Made of' to protein names "
                            "that exist in the Setup protein table")
        analytes.append({"csv_id": p["contents"], "display": p["display"] or p["contents"],
                         "parts": p["parts"], "mw": p["mw"]})
    if not ligands:
        problems.append("no plate column has Role = Ligand")
    if not analytes:
        problems.append("no plate column has Role = Analyte")
    na = exp["n_analytes"]
    if not isinstance(na, (int, float)) or int(na) != na or not 1 <= na <= MAX_ANALYTES:
        problems.append("'Number of analytes' is a manual input: enter a whole number from 1 to "
                        f"{MAX_ANALYTES} (the plate map has {len(analytes)} Analyte column(s))")
    elif int(na) != len(analytes):
        problems.append(f"'Number of analytes' is {int(na)} but {len(analytes)} plate column(s) have "
                        f"Role = Analyte ({[a['csv_id'] for a in analytes]}): fix the number or the roles")
    if len(analytes) > MAX_ANALYTES:
        problems.append(f"{len(analytes)} analytes in the plate map; the limit is {MAX_ANALYTES}")
    if exp["norm_ref"] not in [a["csv_id"] for a in analytes]:
        problems.append(f"Normalization reference analyte {exp['norm_ref']!r} is not one of the "
                        f"analytes {[a['csv_id'] for a in analytes]}")
    start, factor, steps = exp["dil_start"], exp["dil_factor"], exp["dil_steps"]
    if not all(isinstance(x, (int, float)) for x in (start, factor, steps)):
        problems.append("dilution start / factor / steps must be numbers")
        series = []
    else:
        series = [float(start) / float(factor) ** i for i in range(int(steps))]
    if problems:
        raise ValueError(f"{title} sheet is not ready:\n  - " + "\n  - ".join(problems))
    n = exp["n_conc"]
    return {"ligands": ligands, "reference": refs[0] if refs else None, "analytes": analytes,
            "reference_analyte": exp["norm_ref"], "series_M": series,
            "n_concentrations": int(n) if isinstance(n, (int, float)) else None,
            "n_analytes": int(na),
            "results_csv": exp["results_csv"], "prism_template": exp["prism_template"],
            "prism_output": exp["prism_output"],
            "name": exp["name"], "plate": plate}


# ---------------------------------------------------------------------- inference
def _greek(name):
    return name.replace("alpha", "α")


def _match_protein(token, proteins):
    """Best Setup protein name for a plate-entry fragment, or None."""
    t = token.strip().lower()
    t = re.sub(r"\s*complex$", "", t)
    exact = [p for p in proteins if p.lower() == t]
    if exact:
        return exact[0]
    cand = [p for p in proteins if t and t in p.lower() and "(" not in p]
    return min(cand, key=len) if len(cand) >= 1 else None


def infer(setup_path, csv_rows, folder_files, auto=True):
    """First-fill values for the sheet from the Setup sheet, the results CSV and the folder."""
    wf = openpyxl.load_workbook(setup_path)
    wv = openpyxl.load_workbook(setup_path, data_only=True)
    ws = wv[SETUP]
    # protein table: the Proteins tab (A = name, C = MW) or, in a legacy workbook, Setup T / V
    if PROTEINS in wv.sheetnames:
        pws, pn, pm = wv[PROTEINS], 1, 3
    else:
        pws, pn, pm = ws, 20, 22
    loading = {r["Loading Sample ID"] for r in csv_rows}
    samples = [r["Sample ID"] for r in csv_rows]
    ok = [r for r in range(PROT_FIRST, PROT_LAST + 1)
          if isinstance(pws.cell(r, pn).value, str) and isinstance(pws.cell(r, pm).value, (int, float))]
    proteins = [pws.cell(r, pn).value for r in ok]
    mw = {pws.cell(r, pn).value: pws.cell(r, pm).value for r in ok}
    plate, analyte_mw = {}, {}
    entries = scan_plate(ws) if auto else [x for x in scan_plate(ws) if x[1] == 1]   # static: plate row A only
    for idx, (k, first, last, text) in enumerate(entries, 1):
        e = {"col": k}
        if text in loading:
            e["role"], e["part1"] = "Ligand", (_match_protein(text, proteins) or "")
        elif text in samples:
            e["role"] = "Analyte"
            frags = [f for f in text.split("/")]
            ps = [_match_protein(f, proteins) for f in frags]
            e["part1"] = ps[0] or ""
            e["part2"] = ps[1] if len(ps) > 1 and ps[1] else ""
            e["display"] = _greek(text)
            if any(p is None for p in ps):
                e["check"] = "could not match every part to a Setup protein: fill in 'Made of'"
            analyte_mw[text] = sum(mw[p] for p in ps if p)
        elif "glycine" in text.lower():
            e["role"] = "Regeneration"
        elif "hbsa" in text.lower() or "buffer" in text.lower():
            e["role"] = "Buffer"
        elif text in proteins:
            e["role"] = "Reference"       # an immobilized protein that is not in the results
            e["part1"] = text
        if e.get("role") == "Ligand" and not e.get("part1"):
            e["check"] = "no matching protein in the Setup table: fill in 'Made of'"
        plate[idx if auto else k] = e
    csvs = [f for f in folder_files if f.lower().endswith(".csv")]
    prisms = [f for f in folder_files if f.lower().endswith(".prism")]
    biggest = max(analyte_mw, key=analyte_mw.get) if analyte_mw else None
    n_an = sum(1 for e in plate.values() if e.get("role") == "Analyte")
    return {"name": None, "plate": plate, "auto": auto, "analyte_hint": n_an, "n_analytes": None,
            "results_csv": csvs[0] if len(csvs) == 1 else None,
            "prism_template": prisms[0] if len(prisms) == 1 else None,
            "n_conc": None, "norm_ref": biggest, "dil_steps": 8}


def write_with_cache(path, wb):
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    fill_cached_values(path, wb)
