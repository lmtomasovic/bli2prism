"""bli2prism command line entry point.

    draft         write bli2prism_setup.xlsx for a run folder
    build         build the Prism project from the raw files
    workbook      write the Excel workbook (Equilibrium tabs, Kinetic KDs, KD summary)
    qc            write an HTML QC report
    link-proteins make the Proteins tab follow the plate map's roles
    upgrade / fix-contents   repair a setup workbook made by an older version
"""

from __future__ import annotations

import argparse
import sys
import textwrap


def cmd_build(args):
    from .rebuild import build_command
    from .shape import ShapeError
    try:
        return 1 if build_command(args.folder, args.out, args.n_concentrations,
                                  args.template, args.setup, args.no_template, args.blank) else 0
    except ShapeError as e:
        sys.exit(f"cannot build this shape from the template: {e}")
    except (ValueError, FileNotFoundError) as e:
        sys.exit(str(e))


def cmd_workbook(args):
    import os
    from .workbook import build_workbook
    from .rebuild import find_inputs
    out = args.out or os.path.join(args.folder, "rebuild", "workbook.xlsx")
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    try:
        setup, results = find_inputs(args.folder, args.setup)
        res = build_workbook(args.folder, setup, results, out)
    except (ValueError, FileNotFoundError) as e:
        sys.exit(str(e))
    used = sum(1 for c in res["table"] if c["status"] == "used")
    flagged = sum(1 for c in res["table"] if c["flagged"])
    print(f"setup   : {os.path.basename(setup)}")
    print(f"written : {out}")
    print(f"tabs    : {[ws.title for ws in res['wb'].worksheets]}")
    print(f"fit rows: {len(res['table'])} total, {used} used, {flagged} flagged")
    return 0


def cmd_qc(args):
    import os
    from . import qc as Q
    from .rebuild import find_inputs, resolve_prism
    from . import refdata as R
    try:
        setup, results = find_inputs(args.folder, args.setup)
    except (ValueError, FileNotFoundError) as e:
        sys.exit(str(e))
    rows = R.read_results_csv(results)
    n = args.n_concentrations
    if n is None:
        try:
            n = R.read_setup(setup, {r["Loading Sample ID"] for r in rows}).get("n_concentrations")
        except ValueError as e:
            sys.exit(str(e))
    if n is None:
        sys.exit("the number of kinetic concentrations is a manual input: pass -n or fill it in "
                 "on the bli2prism sheet")
    try:
        pick = resolve_prism(args.folder, setup, args.template, args.no_template)
    except (ValueError, FileNotFoundError) as e:
        sys.exit(str(e))
    res = Q.run_qc(args.folder, setup, results, n, pick["template"], scratch=pick["template"] is None,
                   container=pick["container"])
    out = args.out or os.path.join(args.folder, "rebuild", "qc_report.html")
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    Q.write_report(res, out)
    ck = res["checks"]
    print(f"written: {out}")
    print("  ".join(f"{ck.count(l)} {l}" for l in ("FAIL", "WARN", "PASS", "INFO")))
    for c in ck.items:
        if c["level"] in ("FAIL", "WARN"):
            print(f"  {c['level']}: [{c['area']}] {c['title']}")
    return 1 if ck.count("FAIL") else 0


def cmd_upgrade(args):
    import os, shutil
    import openpyxl
    from . import setupsheet as S
    wb = openpyxl.load_workbook(args.file)
    ws = wb[S.sheet_title(wb) or S.SHEET]
    have = {r[0].value for r in ws.iter_rows(min_col=1, max_col=1)}
    missing = [label for _, label, _ in S.KEYS if label not in have]
    backup = os.path.splitext(args.file)[0] + ".before_upgrade.xlsx"
    shutil.copy(args.file, backup)
    try:
        upgraded = S.upgrade_sheet(wb)
    except ValueError as e:
        os.remove(backup)
        raise SystemExit(str(e))
    if upgraded:
        S.write_with_cache(args.file, wb)
        print(f"upgraded {args.file}: added {', '.join(repr(m) for m in missing)}; every input kept\n"
              f"backup  : {backup}")
        if "Number of analytes" in missing:
            print("now enter 'Number of analytes' on the bli2prism sheet")
    else:
        os.remove(backup)
        print("already up to date")
    return 0


def cmd_fix_contents(args):
    import os, shutil
    import openpyxl
    from . import setupsheet as S
    wb = openpyxl.load_workbook(args.file)
    try:
        changes, skipped = S.fix_contents(wb)
    except ValueError as e:
        raise SystemExit(str(e))
    for r, why in skipped:
        print(f"  skipped plate table row {r}: {why}")
    if not changes:
        print("nothing to change: every Ligand / Reference row already reads the name in 'Made of (protein 1)'")
        return 0
    backup = os.path.splitext(args.file)[0] + ".before_fix_contents.xlsx"
    shutil.copy(args.file, backup)
    S.write_with_cache(args.file, wb)
    print(f"re-pointed {len(changes)} Contents formula(s) in {args.file}:")
    for r, old, new in changes:
        print(f"  row {r}: {old}\n       -> {new}")
    print(f"backup  : {backup}")
    return 0


def cmd_link_proteins(args):
    import os, shutil
    import openpyxl
    from . import setupsheet as S
    backup = os.path.splitext(args.file)[0] + ".before_link_proteins.xlsx"
    shutil.copy(args.file, backup)
    wb = openpyxl.load_workbook(args.file)
    done = S.link_proteins(wb, args.immobilized_nM * 1e-9)
    S.write_with_cache(args.file, wb)
    print(f"updated {args.file}:")
    for d in done:
        print(f"  {d}")
    print(f"backup  : {backup}")
    return 0


def cmd_draft(args):
    from .draft import draft
    draft(args.folder, args.source, args.out, auto=not args.static_plate)
    return 0


def main(argv=None):
    p = argparse.ArgumentParser(
        prog="bli2prism",
        description="Octet/BLI raw exports -> GraphPad Prism equilibrium curves and sensorgrams.")
    sub = p.add_subparsers(dest="cmd", required=True)

    bd = sub.add_parser("build", help="build a Prism project from raw files, reshaping a template")
    bd.add_argument("folder", help="folder with the .xlsx, results CSV and A1-H2 .xls files")
    bd.add_argument("-n", "--n-concentrations", type=int,
                    help="kinetics sheets per ligand (dilution steps A..). No default: pass it "
                         "here or fill it in on the bli2prism sheet")
    bd.add_argument("-s", "--setup", help="setup workbook (default: bli2prism_setup.xlsx, else the "
                                          "only .xlsx in FOLDER)")
    bd.add_argument("-t", "--template", help="template .prism (default: the one in FOLDER; with "
                                             "none, the project is built from scratch)")
    bd.add_argument("--no-template", action="store_true",
                    help="build from scratch even if FOLDER has a .prism: data sheets and fit "
                         "analyses, any number of ligands/concentrations, no graphs or layouts")
    bd.add_argument("--blank", help="with --no-template: a blank project saved in Prism, used as "
                                    "the container so your Prism version's skeleton is kept")
    bd.add_argument("-o", "--out", help="output .prism (default: <folder>/rebuild/build.prism)")
    bd.set_defaults(func=cmd_build)

    wk = sub.add_parser("workbook", help="write the Excel workbook (Equilibrium tabs + Kinetic KDs)")
    wk.add_argument("folder")
    wk.add_argument("-s", "--setup", help="workbook holding the Setup sheet (default: the .xlsx in FOLDER)")
    wk.add_argument("-o", "--out", help="output .xlsx (default: <folder>/rebuild/workbook.xlsx)")
    wk.set_defaults(func=cmd_workbook)

    qc = sub.add_parser("qc", help="write an HTML QC report for the run")
    qc.add_argument("folder")
    qc.add_argument("-n", "--n-concentrations", type=int,
                    help="kinetics sheets per ligand (default: from the bli2prism sheet)")
    qc.add_argument("-s", "--setup", help="setup workbook (see build)")
    qc.add_argument("-t", "--template", help="template .prism (default: the one in FOLDER)")
    qc.add_argument("--no-template", action="store_true",
                    help="dry-run the from-scratch build even if FOLDER has a .prism")
    qc.add_argument("-o", "--out", help="output (default: <folder>/rebuild/qc_report.html)")
    qc.set_defaults(func=cmd_qc)

    up = sub.add_parser("upgrade", help="add 'Number of analytes' to an older bli2prism_setup.xlsx "
                                        "(keeps every input; writes a backup first)")
    up.add_argument("file")
    up.set_defaults(func=cmd_upgrade)

    fx = sub.add_parser("fix-contents", help="point wrong Contents formulas of Ligand/Reference rows at the "
                                             "Setup plate-map cells holding the names in 'Made of (protein 1)'")
    fx.add_argument("file")
    fx.set_defaults(func=cmd_fix_contents)

    lp = sub.add_parser("link-proteins", help="make the Proteins tab follow the plate map's roles and add an "
                                              "immobilized-ligand concentration input to Setup")
    lp.add_argument("file")
    lp.add_argument("--immobilized-nM", type=float, default=100.0,
                    help="value written to Setup!R9 if it is empty (default 100 nM)")
    lp.set_defaults(func=cmd_link_proteins)

    dr = sub.add_parser("draft", help="write bli2prism_setup.xlsx: your Setup sheet plus a pre-filled "
                                      "bli2prism sheet inferred from the raw files")
    dr.add_argument("folder")
    dr.add_argument("-s", "--source", help="workbook holding your Setup sheet (default: the only .xlsx)")
    dr.add_argument("-o", "--out", help="output (default: <folder>/bli2prism_setup.xlsx)")
    dr.add_argument("--static-plate", action="store_true",
                    help="one fixed row per plate column (plate row A only) instead of the formula-driven list "
                         "of every plate-map entry")
    dr.set_defaults(func=cmd_draft)

    args = p.parse_args(argv)
    return args.func(args)


if __name__ == "__main__":
    raise SystemExit(main())
