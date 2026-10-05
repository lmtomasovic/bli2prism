"""Excel workbook: one Equilibrium tab per ligand, a Kinetic KDs tab and a KD summary tab.

Equilibrium tabs follow the reference workbook's four-quadrant layout (raw, MW-normalized,
and both normalized to the top concentration = 100), with live formulas.

The Kinetic KDs tab is deliberately NOT the reference layout. It has a summary table above
a per-fit-row table; the summary is formula-driven from the rows' Include column, so
switching a row to No recalculates every average.
"""

from __future__ import annotations

import math

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

from . import eqfit
from . import kinetics as K
from . import refdata as R
from .xlcalc import fill_cached_values

HEAD_FILL = PatternFill("solid", fgColor="DCE6F1")
SECTION_FILL = PatternFill("solid", fgColor="1F3864")
FLAG_FILL = PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid")
BOLD = Font(bold=True)
THIN = Side(style="thin", color="B7B7B7")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

NM_FMT = "[<10]0.000;[<100]0.00;0.0"
SCI_FMT = "0.00E+00"


def _head(ws, ref, text, merge_to=None):
    ws[ref] = text
    c = ws[ref]
    c.font, c.fill, c.alignment, c.border = BOLD, HEAD_FILL, CENTER, BOX
    if merge_to and merge_to != ref:        # a one-cell merge is invalid (one analyte)
        ws.merge_cells(f"{ref}:{merge_to}")


def _sheet_name(prefix, suffixes, used):
    """A unique tab name of at most 31 characters (Excel's limit): prefix + the longest suffix
    that fits; if even the shortest does not, the prefix itself is cut. `suffixes` runs from the
    full wording to the shortest, e.g. [' - Equilibrium', ' - Equil.', ' - Eq.']."""
    name = None
    for suffix in suffixes:
        if len(prefix) + len(suffix) <= 31:
            name = prefix + suffix
            break
    if name is None:
        suffix = suffixes[-1]
        name = prefix[:31 - len(suffix)] + suffix
    base, n = name, 2
    while name in used:
        tag = f" ({n})"
        name = base[:31 - len(tag)] + tag
        n += 1
    used.add(name)
    return name


# -------------------------------------------------------------------- equilibrium
def write_equilibrium(ws, ligand, reference, xs_M, columns, mw_cells_spec, t_assoc, ref_idx):
    """columns: list over analytes of (display_name, [response per concentration]).

    mw_cells_spec: list over analytes of either a float (literal MW) or a list of analyte
    indexes whose MWs are summed.
    """
    n = len(xs_M)
    na = len(columns)
    raw = [4 + i for i in range(na)]                  # D, E, F ...
    norm = [4 + na + 1 + i for i in range(na)]        # H, I, J ...
    ws["A1"] = (f"Background subtracted from {reference} loaded biosensor"
                if reference else "Background-subtracted data")
    ws["A1"].font = Font(italic=True)

    def block(top, title_raw, title_norm, sub, sub_norm, first_data, is_normalized100):
        _head(ws, f"{L(raw[0])}{top}", title_raw, f"{L(raw[-1])}{top}")
        _head(ws, f"{L(norm[0])}{top}", title_norm, f"{L(norm[-1])}{top}")
        _head(ws, f"{L(raw[0])}{top + 1}", sub, f"{L(raw[-1])}{top + 1}")
        _head(ws, f"{L(norm[0])}{top + 1}", sub_norm, f"{L(norm[-1])}{top + 1}")
        for i, (nm, _) in enumerate(columns):
            _head(ws, f"{L(raw[i])}{top + 2}", nm)
            _head(ws, f"{L(norm[i])}{top + 2}", nm)
        mw_row = top + 3
        for i in range(na):
            spec = mw_cells_spec[i]
            if top == 2:      # the MW source row
                if isinstance(spec, list):
                    ws[f"{L(raw[i])}{mw_row}"] = "=" + "+".join(f"{L(raw[j])}{mw_row}" for j in spec)
                else:
                    ws[f"{L(raw[i])}{mw_row}"] = spec
            else:
                ws[f"{L(raw[i])}{mw_row}"] = f"={L(raw[i])}$5"
            ws[f"{L(norm[i])}{mw_row}"] = f"={L(raw[i])}$5"
            for col in (raw[i], norm[i]):
                ws[f"{L(col)}{mw_row}"].number_format = "#,##0.00"
        ws[f"{L(raw[0] - 1)}{mw_row}"] = "MW (g/mol)"
        ws[f"{L(raw[0] - 1)}{mw_row}"].font = BOLD
        lab = top + 4
        for col, txt in ((1, "[concentration] (M)"), (2, "[concentration] (nM)"), (3, "log")):
            _head(ws, f"{L(col)}{lab}", txt)
        if not is_normalized100:
            _head(ws, f"{L(raw[0])}{lab}", f"Response at X = {t_assoc} s (nm)", f"{L(raw[-1])}{lab}")
            _head(ws, f"{L(norm[0])}{lab}", f"MW-normalized to {ref_name}", f"{L(norm[-1])}{lab}")
        return first_data

    ref_name = columns[ref_idx][0]
    # --- block 1: raw + MW normalized
    d0 = 7
    block(2, "Raw Data", "MW Normalized", "Raw Binding (nm)", "Normalized Binding (nm)", d0, False)
    ws[f"{L(norm[-1] + 1)}4"] = "Reference MW"
    ws[f"{L(norm[-1] + 1)}4"].font = BOLD
    ws[f"{L(norm[-1] + 1)}5"] = f"={L(norm[ref_idx])}5"
    ws[f"{L(norm[-1] + 1)}5"].number_format = "#,##0.00"
    refcell = f"${L(norm[-1] + 1)}$5"
    for r in range(n):
        row = d0 + r
        ws[f"A{row}"] = xs_M[r]
        ws[f"B{row}"] = f"=A{row}*10^9"
        ws[f"C{row}"] = f"=LOG(A{row})"
        for i in range(na):
            ws[f"{L(raw[i])}{row}"] = columns[i][1][r]
            ws[f"{L(norm[i])}{row}"] = f"={L(raw[i])}{row}*({refcell}/{L(norm[i])}$5)"
    # --- block 2: normalized to 100
    top2 = d0 + n + 1
    d2 = top2 + 5
    block(top2, "Normalized to 100", "MW Normalized to 100", "Normalized Binding (nm)",
          "Normalized Binding (nm)", d2, True)
    for r in range(n):
        row = d2 + r
        ws[f"A{row}"] = f"=A{d0 + r}"
        ws[f"B{row}"] = f"=B{d0 + r}"
        ws[f"C{row}"] = f"=C{d0 + r}"
        for i in range(na):
            ws[f"{L(raw[i])}{row}"] = f"={L(raw[i])}{d0 + r}/{L(raw[i])}${d0}*100"
            ws[f"{L(norm[i])}{row}"] = f"={L(norm[i])}{d0 + r}/{L(norm[i])}${d0}*100"
    # formats
    for rows in (range(d0, d0 + n), range(d2, d2 + n)):
        for row in rows:
            ws[f"A{row}"].number_format = SCI_FMT
            ws[f"B{row}"].number_format = "0.0###"
            ws[f"C{row}"].number_format = "0.000"
            for col in raw + norm:
                ws[f"{L(col)}{row}"].number_format = "0.0000" if rows.start == d0 else "0.0"
                ws[f"{L(col)}{row}"].border = BOX
            for col in (1, 2, 3):
                ws[f"{L(col)}{row}"].border = BOX
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 10
    for col in raw + norm:
        ws.column_dimensions[L(col)].width = 17
    ws.column_dimensions[L(raw[-1] + 1)].width = 3
    ws.column_dimensions[L(norm[-1] + 1)].width = 15
    ws.sheet_view.showGridLines = False


# ----------------------------------------------------------------------- kinetic KDs
def write_kinetic_kds(ws, table, ligands, analytes):
    ws["A1"] = "Kinetic KDs"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (f"Fit rows used: Full X² < {K.X2_MAX:g}, Full R² > {K.R2_MIN:g}, and KD, kon and "
                "kdis all determinate (not censored or non-converged). Only the KD column is used; "
                "the 2:1 second-site KD2 is ignored.")
    ws["A3"] = ("Flagged = used, but KD is more than "
                f"{K.OUTLIER_FOLD:g}-fold from the median of the group's used rows; flagged "
                "rows stay included. Change a row's Include to No in the table below and the "
                "summary recalculates.")
    for r in (2, 3):
        ws[f"A{r}"].font = Font(italic=True, color="595959")

    n_groups = len(ligands) * len(analytes)
    sum_head = 5
    sum0 = sum_head + 2
    sum1 = sum0 + n_groups - 1
    det_title = sum1 + 4
    det_head = det_title + 1
    det0 = det_head + 1
    det1 = det0 + len(table) - 1

    # ---- summary
    ws[f"A{sum_head}"] = "Summary"
    ws[f"A{sum_head}"].font = Font(bold=True, size=12)
    heads = ["Ligand", "Analyte", "Rows used", "Rows flagged", "Mean kon (1/Ms)",
             "Mean kdis (1/s)", "KD from mean kdis / mean kon (nM)", "Mean of per-row KDs (nM)",
             "Ratio of the two KDs", "Agreement", "Mean KD excl. flagged rows (nM)", "Notes"]
    for i, h in enumerate(heads, 1):
        _head(ws, f"{L(i)}{sum_head + 1}", h)
    ws.row_dimensions[sum_head + 1].height = 48

    def rng(col):
        return f"${col}${det0}:${col}${det1}"

    r = sum0
    for lig in ligands:
        for a in analytes:
            crit = f'{rng("A")},$A{r},{rng("B")},$B{r},{rng("D")},"Yes"'
            ws[f"A{r}"], ws[f"B{r}"] = lig, a
            ws[f"C{r}"] = f"=COUNTIFS({crit})"
            ws[f"D{r}"] = f'=COUNTIFS({crit},{rng("K")},"Yes")'
            ws[f"E{r}"] = f'=IFERROR(AVERAGEIFS({rng("F")},{crit}),"n/a")'
            ws[f"F{r}"] = f'=IFERROR(AVERAGEIFS({rng("G")},{crit}),"n/a")'
            ws[f"G{r}"] = f'=IFERROR(F{r}/E{r}*10^9,"n/a")'
            ws[f"H{r}"] = f'=IFERROR(AVERAGEIFS({rng("E")},{crit}),"n/a")'
            ws[f"I{r}"] = f'=IFERROR(MAX(G{r},H{r})/MIN(G{r},H{r}),"n/a")'
            ws[f"J{r}"] = f'=IF(ISNUMBER(I{r}),IF(I{r}>2,"differ","agree"),"")'
            ws[f"K{r}"] = f'=IFERROR(AVERAGEIFS({rng("E")},{crit},{rng("K")},"<>Yes"),"n/a")'
            grp = [c for c in table if c["ligand"] == lig and c["analyte"] == a]
            used = [c for c in grp if c["status"] == "used"]
            note = []
            if not used:
                note.append("No fit row passed the filters.")
            only2 = [c for c in grp if "second site" in c["reason"]]
            if only2:
                note.append(f"{len(only2)} row(s) have a censored KD but a determinate 2:1 "
                            "second-site KD2 (see Note column below); KD2 is not used.")
            ws[f"L{r}"] = " ".join(note)
            for col in range(1, 13):
                c = ws[f"{L(col)}{r}"]
                c.border = BOX
                c.alignment = Alignment(vertical="center", wrap_text=(col == 12),
                                        horizontal="left" if col in (1, 2, 12) else "center")
            for col in "EF":
                ws[f"{col}{r}"].number_format = SCI_FMT
            for col in "GHK":
                ws[f"{col}{r}"].number_format = NM_FMT
            ws[f"I{r}"].number_format = "0.0"
            r += 1
    ws.conditional_formatting.add(f"J{sum0}:J{sum1}", FormulaRule(
        formula=[f'$J{sum0}="differ"'], fill=FLAG_FILL))
    ws.conditional_formatting.add(f"D{sum0}:D{sum1}", FormulaRule(
        formula=[f"$D{sum0}>0"], fill=FLAG_FILL))
    ws[f"A{sum1 + 1}"] = ("Ratio = larger / smaller of the two KD estimates; 'differ' when above 2. "
                          "'n/a' = no included rows.")
    ws[f"A{sum1 + 1}"].font = Font(italic=True, color="595959", size=9)

    # ---- detail
    ws[f"A{det_title}"] = "Fit rows"
    ws[f"A{det_title}"].font = Font(bold=True, size=12)
    dheads = ["Ligand", "Analyte", "Conc (nM)", "Include", "KD (nM)", "kon (1/Ms)", "kdis (1/s)",
              "Full X²", "Full R²", "Note", "Flagged"]
    for i, h in enumerate(dheads, 1):
        _head(ws, f"{L(i)}{det_head}", h)
    for k, c in enumerate(table):
        row = det0 + k
        ws[f"A{row}"], ws[f"B{row}"] = c["ligand"], c["analyte"]
        ws[f"C{row}"] = c["conc_nM"]
        ws[f"D{row}"] = "Yes" if c["status"] == "used" else "No"
        ws[f"E{row}"] = c["kd_M"] * 1e9 if c["kd_M"] is not None else f"{c['kd_text']} M"
        ws[f"F{row}"] = c["kon"] if c["kon"] is not None and c["kon"] < K.NONCONVERGED else c["kon_text"]
        ws[f"G{row}"] = c["kdis"] if c["kdis"] is not None and c["kdis"] < K.NONCONVERGED else c["kdis_text"]
        ws[f"H{row}"], ws[f"I{row}"] = c["x2"], c["r2"]
        ws[f"J{row}"] = c["reason"]
        ws[f"K{row}"] = "Yes" if c["flagged"] else None
        ws[f"C{row}"].number_format = "0.0##"
        ws[f"E{row}"].number_format = NM_FMT
        for col in "FG":
            ws[f"{col}{row}"].number_format = SCI_FMT
        for col in "HI":
            ws[f"{col}{row}"].number_format = "0.000"
        for col in range(1, 12):
            cell = ws[f"{L(col)}{row}"]
            cell.border = BOX
            cell.alignment = Alignment(vertical="center",
                                       horizontal="left" if col in (1, 2, 10) else "center")
    dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"D{det0}:D{det1}")
    ws.conditional_formatting.add(f"A{det0}:K{det1}", FormulaRule(
        formula=[f'$D{det0}="No"'], font=Font(color="9A9A9A")))
    ws.conditional_formatting.add(f"A{det0}:K{det1}", FormulaRule(
        formula=[f'$K{det0}="Yes"'], fill=FLAG_FILL))

    widths = [18, 28, 12, 12, 16, 16, 20, 18, 14, 46, 16, 60]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[L(i)].width = w
    ws.sheet_view.showGridLines = False
    return {"summary": (sum0, sum1), "detail": (det0, det1)}


def _nM(log10_M):
    """10^log10(M) in nM, or 'n/a' when the value is missing or too large to be a number (a fit to a trace that never
    binds can put an interval end at 1e300 or beyond)."""
    try:
        v = 10 ** log10_M * 1e9
    except OverflowError:
        return "n/a"
    return v if math.isfinite(v) and v < 1e15 else "n/a"


def write_kd_summary(ws, ligands, analytes, fits, kd_sheet, kd_rows):
    """One row per ligand / analyte: equilibrium KD (nM) next to the mean per-row kinetic KD (nM).

    The equilibrium KD is the EC50 of the sigmoidal fit that the Prism analyses apply (fixed slope,
    Y = Bottom + (Top - Bottom) / (1 + 10^(LogEC50 - X)), X = log10 M), computed here with the same
    model; Prism refits on open, so its own EC50 should agree to fit precision. The kinetic column is a
    live link to the Kinetic KDs tab's 'Mean of per-row KDs', so switching a row's Include there updates it.
    fits: {(ligand, analyte display name): eqfit.fit_curve result}; kd_rows: {(ligand, analyte): row on kd_sheet}.
    """
    ws["A1"] = "KD summary (nM)"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = ("Equilibrium KD = EC50 of the sigmoidal equilibrium fit (the fit the Prism analyses apply), converted "
                "from log10 M to nM. Kinetic KD = mean of the per-row KDs on the Kinetic KDs tab (included rows only).")
    ws["A3"] = ("The equilibrium value is calculated here with the same model Prism uses, not read from Prism's own results "
                "table; check it against the Prism analysis. Intervals are approximate (asymptotic, Prism reports profile-likelihood).")
    for r in (2, 3):
        ws[f"A{r}"].font = Font(italic=True, color="595959")
    heads = ["Ligand", "Analyte", "Equilibrium KD (nM)", "95% CI low (nM)", "95% CI high (nM)", "Fit R²",
             "Mean per-row kinetic KD (nM)", "Kinetic rows used", "Equilibrium notes"]
    for i, h in enumerate(heads, 1):
        _head(ws, f"{L(i)}5", h)
    ws.row_dimensions[5].height = 48
    q = "'" + kd_sheet.replace("'", "''") + "'"
    r = 6
    for lig in ligands:
        for a in analytes:
            f = fits[(lig, a)]
            t = f["three"]
            ws[f"A{r}"], ws[f"B{r}"] = lig, a
            if t["ok"]:
                lec = t["params"][2]
                lo, hi = t["ci"][2]
                ws[f"C{r}"] = _nM(lec)
                ws[f"D{r}"] = _nM(lo)
                ws[f"E{r}"] = _nM(hi)
                ws[f"F{r}"] = t["r2"]
            else:
                for col in "CDEF":
                    ws[f"{col}{r}"] = "n/a"
            ws[f"G{r}"] = f"={q}!H{kd_rows[(lig, a)]}"
            ws[f"H{r}"] = f"={q}!C{kd_rows[(lig, a)]}"
            ws[f"I{r}"] = "; ".join(eqfit.flags(f))
            for col in range(1, 10):
                c = ws[f"{L(col)}{r}"]
                c.border = BOX
                c.alignment = Alignment(vertical="center", wrap_text=(col == 9),
                                        horizontal="left" if col in (1, 2, 9) else "center")
            for col in "CDEG":
                ws[f"{col}{r}"].number_format = NM_FMT
            ws[f"F{r}"].number_format = "0.000"
            r += 1
    ws.conditional_formatting.add(f"I6:I{r - 1}", FormulaRule(formula=['$I6<>""'], fill=FLAG_FILL))
    for i, w in enumerate([18, 28, 18, 16, 16, 10, 22, 12, 70], 1):
        ws.column_dimensions[L(i)].width = w
    ws.sheet_view.showGridLines = False


# ------------------------------------------------------------------------- driver
def analyte_mw(analytes):
    """Per analyte: its MW (float), or the list of analyte indexes whose MWs it is the sum of
    (the complex = its two components, as in the reference workbook's =SUM(E5:F5))."""
    singles = {a["parts"][0]: i for i, a in enumerate(analytes) if len(a["parts"]) == 1}
    spec = []
    for a in analytes:
        if len(a["parts"]) > 1 and all(p in singles for p in a["parts"]):
            spec.append([singles[p] for p in a["parts"]])
        else:
            spec.append(float(a["mw"]))
    return spec


def build_workbook(folder, setup_xlsx, results_csv, out):
    csv_rows = R.read_results_csv(results_csv)
    setup = R.read_setup(setup_xlsx, {r["Loading Sample ID"] for r in csv_rows})
    ligands, analytes = setup["ligands"], setup["analytes"]
    analytes_csv = [a["csv_id"] for a in analytes]
    display = {a["csv_id"]: a["display"] for a in analytes}
    assoc = [k for k in csv_rows[0] if k.startswith("X=") and k != "X=0"][0]
    t_assoc = assoc[2:]
    mw_spec = analyte_mw(analytes)
    ref_idx = analytes_csv.index(setup["reference_analyte"])

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    used = set()
    eq_fits = {}
    files_by_ligand = R.ligand_files(csv_rows, ligands, setup["series_M"])
    for lig in ligands:
        files = files_by_ligand[lig]
        present = sorted(files)
        by_step, warns = R.step_concentrations(folder, files, setup["series_M"], True)
        for w in warns:
            print("WARNING:", w)
        xs = [by_step[s_] for s_ in present]
        lines = R.equilibrium_table(csv_rows, lig, setup["series_M"], xs, analytes, steps=present)
        cols = [(display[a], [float(l.split(",")[1 + j]) for l in lines])
                for j, a in enumerate(analytes_csv)]
        for (disp, y) in cols:
            eq_fits[(lig, disp)] = eqfit.fit_curve([math.log10(x) for x in xs], y)
        ws = wb.create_sheet(_sheet_name(lig, [" - Equilibrium", " - Equil.", " - Eq."], used))
        write_equilibrium(ws, lig, setup["reference"], xs, cols, mw_spec, t_assoc, ref_idx)
    fit_rows = K.read_fit_rows(results_csv)
    table = K.build_table(fit_rows, ligands, analytes_csv)
    for c in table:
        c["analyte"] = display[c["analyte"]]
    ws = wb.create_sheet(_sheet_name("Kinetic KDs", [""], used))
    layout = write_kinetic_kds(ws, table, ligands, [display[a] for a in analytes_csv])
    disp_list = [display[a] for a in analytes_csv]
    kd_rows = {(lig, a): layout["summary"][0] + i * len(disp_list) + j
               for i, lig in enumerate(ligands) for j, a in enumerate(disp_list)}
    ws_sum = wb.create_sheet(_sheet_name("KD summary", [""], used))
    write_kd_summary(ws_sum, ligands, disp_list, eq_fits, ws.title, kd_rows)
    wb.calculation.fullCalcOnLoad = True
    wb.save(out)
    values = fill_cached_values(out, wb)
    return {"ligands": ligands, "table": table, "layout": layout, "values": values,
            "kd_sheet": ws.title, "wb": wb}
