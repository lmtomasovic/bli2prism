"""The auto-listing plate table: every Setup plate-map entry is listed by formula, whatever the layout."""

import openpyxl
import pytest

from bli2prism import setupsheet as S
from bli2prism import xlcalc as X


def make_wb(cells):
    """cells: {'B5': 'name', ...} on a Setup sheet; Proteins tab with a few MWs."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = S.SETUP
    ws["N5"], ws["R5"] = 1e-6, 3
    for ref, v in cells.items():
        ws[ref] = v
    p = wb.create_sheet(S.PROTEINS)
    for i, (n, mw) in enumerate([("Lig", 50000), ("Ref", 40000), ("An", 10000)]):
        p[f"A{5 + i}"], p[f"C{5 + i}"] = n, mw
    return wb


def listed(wb):
    """(col, name, rows, label) per table row, as Excel would compute them (via xlcalc)."""
    ws = wb[S.SHEET]
    ev = X.Evaluator(wb)
    labels = {r[0].value: r[0].row for r in ws.iter_rows(min_col=1, max_col=1) if isinstance(r[0].value, str)}
    body, _ = S.plate_rows(ws, labels["[Plate columns]"])
    out = [tuple(ev.value(f"{c}{r}", S.SHEET) for c in "ABIJ") for r in body]
    return [o for o in out if o[1] != ""], ws, ev


def build(cells, roles=None):
    wb = make_wb(cells)
    entries = S.scan_plate(wb[S.SETUP])
    plate = {i: {"col": k, "role": (roles or {}).get(t)} for i, (k, a, b, t) in enumerate(entries, 1)}
    ws = S.add_sheet(wb, {"plate": plate, "auto": True})
    S.style_setup(wb[S.SETUP], auto_rows=S.auto_row_count(ws))
    return wb, entries


# 1. whole column of reference sensors; 2. one reference sensor above seven ligand sensors in a column;
# 3. two ligands stacked A-D / E-H; 4. a name repeated down a column of merged cells.
LAYOUTS = {
    "reference column": {"B5": "Lig", "C5": "Ref", "D5": "An", "E5": "Buffer"},
    "single reference in the ligand column": dict(
        {"B5": "Ref"}, **{f"B{r}": "Lig" for r in range(6, 13)}, **{"C5": "An"}),
    "stacked ligands": {"B5": "Lig A", "B9": "Lig B", "C5": "An"},
    "repeated names": {f"B{r}": "Lig" for r in range(5, 13)},
}


@pytest.mark.parametrize("name", list(LAYOUTS))
def test_formulas_list_what_scan_plate_lists(name):
    wb, entries = build(LAYOUTS[name])
    got, _, _ = listed(wb)
    want = [(k, t, "ABCDEFGH"[a - 1] if a == b else f"{'ABCDEFGH'[a - 1]}-{'ABCDEFGH'[b - 1]}")
            for k, a, b, t in entries]
    assert [g[:3] for g in got] == want


def test_single_reference_sensor_in_a_ligand_column():
    wb, _ = build(LAYOUTS["single reference in the ligand column"])
    got, _, _ = listed(wb)
    assert got[0][:3] == (1, "Ref", "A") and got[1][:3] == (1, "Lig", "B-H")      # 7 ligand sensors


def test_stacked_ligands_get_their_own_row_ranges():
    wb, _ = build(LAYOUTS["stacked ligands"])
    got, _, _ = listed(wb)
    assert [g[:3] for g in got] == [(1, "Lig A", "A-D"), (1, "Lig B", "E-H"), (2, "An", "A-H")]


def test_repeated_name_is_one_entry():
    wb, _ = build(LAYOUTS["repeated names"])
    got, _, _ = listed(wb)
    assert [g[:3] for g in got] == [(1, "Lig", "A-H")]


def test_plate_map_edit_changes_the_list_without_touching_the_sheet():
    wb, _ = build(LAYOUTS["reference column"])
    wb[S.SETUP]["F5"] = "Extra"
    got, _, _ = listed(wb)
    assert (5, "Extra", "A-H") in [g[:3] for g in got]


def test_setup_role_row_shows_mixed_columns():
    wb, _ = build(LAYOUTS["single reference in the ligand column"],
                  roles={"Ref": "Reference", "Lig": "Ligand", "An": "Analyte"})
    ws = wb[S.SHEET]
    p0 = S.plate_row0()
    for i, role in enumerate(("Reference", "Ligand", "Analyte")):
        ws[f"C{p0 + i}"] = role
    ev = X.Evaluator(wb)
    row3 = [ev.value(f"{c}3", S.SETUP) for c in "BCD"]
    assert row3 == ["Reference +1 more", "Analyte 1", ""]


def test_reader_accepts_an_auto_sheet(tmp_path):
    wb, _ = build(LAYOUTS["single reference in the ligand column"])
    ws = wb[S.SHEET]
    p0 = S.plate_row0()
    for i, (role, part) in enumerate((("Reference", "Ref"), ("Ligand", "Lig"), ("Analyte", "An"))):
        ws[f"C{p0 + i}"], ws[f"E{p0 + i}"] = role, part
    ws[f"D{p0 + 2}"] = "An"
    labels = {ws[f"A{r}"].value: r for r in range(5, 20) if ws[f"A{r}"].value}
    for label, v in (("Number of analytes", 1), ("Normalization reference analyte", "An"),
                     ("Dilution steps", 4), ("Results CSV", "x.csv")):
        ws[f"B{labels[label]}"] = v
    path = str(tmp_path / "s.xlsx")
    S.write_with_cache(path, wb)
    s = S.read(path)
    assert s["ligands"] == ["Lig"] and s["reference"] == "Ref"
    assert [a["csv_id"] for a in s["analytes"]] == ["An"] and s["analytes"][0]["mw"] == 10000
    assert [e["col"] for e in s["plate"] if e["contents"]] == [1, 1, 2]


def test_repair_commands_refuse_an_auto_sheet():
    wb, _ = build(LAYOUTS["reference column"])
    with pytest.raises(ValueError, match="nothing to fix"):
        S.fix_contents(wb)
    with pytest.raises(ValueError, match="cannot be upgraded"):
        wb[S.SHEET]["A5"] = "stale"
        S.upgrade_sheet(wb)


@pytest.mark.parametrize("auto", [True, False])
def test_protein_and_display_dropdowns(auto):
    wb = make_wb(LAYOUTS["reference column"])
    spec = {"plate": {1: {"col": 1}, 2: {"col": 2}}, "auto": auto}
    ws = S.add_sheet(wb, spec)
    p0 = S.plate_row0()
    dvs = {str(next(iter(d.sqref.ranges)).coord): d for d in ws.data_validations.dataValidation}
    n = S.auto_row_count(ws)
    prot = dvs[f"E{p0}:F{p0 + n - 1}"]
    assert prot.formula1 == f"={S.PROTEIN_NAMES}" and prot.showErrorMessage
    assert wb.defined_names[S.PROTEIN_NAMES].attr_text == f"{S.PROTEINS}!$A${S.PROT_FIRST}:$A${S.PROT_LAST}"
    disp = dvs[f"D{p0}:D{p0 + n - 1}"]
    assert disp.formula1 == f"=$B${p0}:$B${p0 + n - 1}" and not disp.showErrorMessage


def test_draft_accepts_a_dated_csv_name_and_an_existing_setup_as_source(demo_run):
    import contextlib, glob, io, os, shutil
    from bli2prism.draft import draft
    d = demo_run
    src = os.path.join(d, "bli2prism_setup.xlsx")
    first = os.path.join(d, "first.xlsx")
    with contextlib.redirect_stdout(io.StringIO()):
        draft(d, source=src, out=first)
        draft(d, source=first, out=os.path.join(d, "second.xlsx"))      # source already has a Proteins tab
    wb = openpyxl.load_workbook(os.path.join(d, "second.xlsx"))
    assert wb.sheetnames[:1] == ["bli2prism"] and S.PROTEINS in wb.sheetnames
    assert wb[S.PROTEINS]["A5"].value == openpyxl.load_workbook(first)[S.PROTEINS]["A5"].value
    for f in glob.glob(os.path.join(d, "kineticanalysistableresults*.csv")):
        shutil.copy(f, os.path.join(d, "kineticanalysistableresults_extra.csv"))
    with pytest.raises(SystemExit, match="exactly one"):
        draft(d, source=first, out=os.path.join(d, "third.xlsx"))


# ------------------------------------------------------------ Proteins tab follows the plate map's roles
def protein_wb(extra_role=None):
    wb, _ = build(LAYOUTS["single reference in the ligand column"])
    ws = wb[S.SHEET]
    p0 = S.plate_row0()
    for i, (role, part) in enumerate((("Reference", "Ref"), ("Ligand", "Lig"), ("Analyte", "An"))):
        ws[f"C{p0 + i}"], ws[f"E{p0 + i}"] = role, part
    pr = wb[S.PROTEINS]
    pr["A8"], pr["C8"] = "Unused", 1000
    for r in range(5, 9):                       # the lab's table: formulas in E.., one row pointing one row up
        pr[f"E{r}"] = f"=B{r}/D{r}"
        pr[f"G{r}"] = f"=H{r - 1}*I{r - 1}*10^-6"
        pr[f"I{r}"] = 2000
    return wb


def test_link_proteins_sets_desired_concentration_by_role():
    wb = protein_wb()
    wb["Setup"]["R9"] = 2e-7                     # user's immobilized concentration
    S.link_proteins(wb)
    ev = X.Evaluator(wb)
    got = {wb["Proteins"][f"A{r}"].value: (ev.value(f"N{r}", "Proteins"), ev.value(f"H{r}", "Proteins")) for r in range(5, 9)}
    assert got["Lig"] == ("Ligand", 2e-7) and got["Ref"] == ("Reference", 2e-7)
    assert got["An"] == ("Analyte", 1e-6)                                  # Setup!N5, the starting concentration
    assert got["Unused"] == ("", 0)


def test_link_proteins_follows_edits_to_the_input_and_the_roles():
    wb = protein_wb()
    S.link_proteins(wb)
    assert wb["Setup"]["R9"].value == 1e-7 and wb["Setup"]["Q9"].value == "Immobilized ligand concentration (M)"
    wb["Setup"]["R9"] = 5e-8
    ws = wb[S.SHEET]
    ws[f"C{S.plate_row0()}"], ws[f"E{S.plate_row0()}"] = "Ligand", "An"      # An is now also loaded as a ligand
    ev = X.Evaluator(wb)
    assert ev.value("N7", "Proteins") == "Ligand + Analyte" and ev.value("H7", "Proteins") == 1e-6
    assert ev.value("H5", "Proteins") == 5e-8


def test_link_proteins_repoints_moles_needed_and_keeps_totals_working():
    wb = protein_wb()
    S.link_proteins(wb)
    assert [wb["Proteins"][f"G{r}"].value for r in range(5, 9)] == [f"=H{r}*I{r}*10^-6" for r in range(5, 9)]
    ev = X.Evaluator(wb)
    assert ev.value("G6", "Proteins") == pytest.approx(ev.value("H6", "Proteins") * ev.value("I6", "Proteins") * 1e-6)


def test_link_proteins_reasserts_the_made_of_dropdowns_and_is_idempotent():
    wb = protein_wb()
    S.link_proteins(wb)
    S.link_proteins(wb)
    ws = wb[S.SHEET]
    p0 = S.plate_row0()
    lists = [(str(d.sqref), d.formula1) for d in ws.data_validations.dataValidation if d.formula1 == f"={S.PROTEIN_NAMES}"]
    assert len(lists) == 1 and lists[0][0].startswith(f"E{p0}:F")
    assert not any("E" in str(d.sqref) and d.formula1.startswith("=$B$") for d in ws.data_validations.dataValidation
                   if f"E{p0}" in str(d.sqref))
    assert wb["Setup"]["R9"].value == 1e-7


def test_link_proteins_cli_backs_up_and_survives_a_save(tmp_path):
    import subprocess
    wb = protein_wb()
    path = str(tmp_path / "s.xlsx")
    wb.save(path)
    from conftest import run_cli
    run_cli("link-proteins", path)
    assert (tmp_path / "s.before_link_proteins.xlsx").exists()
    again = openpyxl.load_workbook(path)
    assert again["Proteins"]["N6"].value.startswith("=IF(") and again["Setup"]["R9"].value == pytest.approx(1e-7)


def test_total_volume_follows_the_role():
    wb = protein_wb()
    wb["Setup"]["R8"] = 285                       # Final Volume + Transfer Volume
    S.link_proteins(wb)
    ev = X.Evaluator(wb)
    vol = {wb["Proteins"][f"A{r}"].value: ev.value(f"I{r}", "Proteins") for r in range(5, 9)}
    assert vol["An"] == pytest.approx(285 * 1.2)                          # analyte: (final + transfer volume) x 1.2
    assert vol["Lig"] == pytest.approx(7 * 200 * 1.2)                     # ligand: 7 sensor rows (B-H) x 200 uL x 1.2
    assert vol["Ref"] == pytest.approx(1 * 200 * 1.2)                     # the single reference sensor
    assert vol["Unused"] == 0
    wb["Setup"]["R10"], wb["Setup"]["R11"] = 150, 1.5                      # the two constants are inputs
    ev = X.Evaluator(wb)
    assert ev.value("I5", "Proteins") == pytest.approx(7 * 150 * 1.5) and ev.value("I7", "Proteins") == pytest.approx(285 * 1.5)


def test_sensor_rows_count_matches_the_listed_row_ranges():
    for name in ("stacked ligands", "reference column", "single reference in the ligand column"):
        wb, entries = build(LAYOUTS[name])
        ws = wb[S.SHEET]
        ev = X.Evaluator(wb)
        p0 = S.plate_row0()
        got = [ev.value(f"K{p0 + i}", S.SHEET) for i in range(len(entries))]
        assert got == [b - a + 1 for _, a, b, _ in entries]
        assert ws[f"K{p0 - 1}"].value == S.SENSOR_ROWS_HEAD


def test_link_proteins_adds_the_count_column_to_an_older_auto_sheet():
    wb = protein_wb()
    ws = wb[S.SHEET]
    p0 = S.plate_row0()
    for r in range(p0 - 1, p0 + 14):                   # an auto sheet drafted before column K existed
        ws[f"K{r}"] = None
    wb["Setup"]["R8"] = 285
    S.link_proteins(wb)
    ev = X.Evaluator(wb)
    assert ws[f"K{p0 - 1}"].value == S.SENSOR_ROWS_HEAD
    assert ev.value("I5", "Proteins") == pytest.approx(7 * 200 * 1.2)
