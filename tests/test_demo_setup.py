"""Setup workbook, renamed tabs, never overwriting, and the workbook's KD summary, on the demo experiment."""

import os

import openpyxl
import pytest

from bli2prism import setupsheet as S
from bli2prism import xlcalc as X
from conftest import run_cli


def setup_path(d):
    return os.path.join(d, "bli2prism_setup.xlsx")


def test_the_demo_setup_reads(demo_run):
    s = S.read(setup_path(demo_run))
    assert s["ligands"] == ["Binder A", "Binder B"] and s["reference"] == "Control IgG"
    assert [a["csv_id"] for a in s["analytes"]] == ["Analyte X", "Analyte Y", "Analyte Z"]
    mw = {a["csv_id"]: a["mw"] for a in s["analytes"]}
    assert mw == {"Analyte X": 21000, "Analyte Y": 22000, "Analyte Z": 25000}
    assert s["n_concentrations"] == 8 and s["n_analytes"] == 3
    assert s["series_M"][:3] == pytest.approx([1e-6, 1e-6 / 3, 1e-6 / 9])


def test_a_wrong_analyte_count_is_explained(demo_run):
    wb = openpyxl.load_workbook(setup_path(demo_run))
    ws = wb[S.SHEET]
    row = next(r for r in range(5, 16) if ws[f"A{r}"].value == "Number of analytes")
    ws[f"B{row}"] = 2
    S.write_with_cache(setup_path(demo_run), wb)
    with pytest.raises(ValueError, match="3 plate column"):
        S.read(setup_path(demo_run))


RENAMES = {"bli2prism": "Inputs", "Setup": "Plate Map", "Proteins": "Lab Proteins"}


def rename_tabs(wb):
    """Rename the tabs the way Excel does: formulas that point at a tab follow its new name."""
    names = {old: S.q(new) for old, new in RENAMES.items()}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    for old, new in names.items():
                        c.value = c.value.replace(f"{old}!", f"{new}!")
    for old, new in RENAMES.items():
        wb[old].title = new
    return wb


def test_renamed_tabs_are_found_and_read(demo_run):
    wb = rename_tabs(openpyxl.load_workbook(setup_path(demo_run)))
    assert S.detect_names(wb) == ("Inputs", "Plate Map", "Lab Proteins")
    S.write_with_cache(setup_path(demo_run), wb)
    assert S.has_sheet(setup_path(demo_run)) and S.read(setup_path(demo_run))["ligands"] == ["Binder A", "Binder B"]


def test_link_proteins_works_on_renamed_tabs(demo_run):
    wb = rename_tabs(openpyxl.load_workbook(setup_path(demo_run)))
    S.link_proteins(wb)
    pr = wb["Lab Proteins"]
    ev = X.Evaluator(wb)
    role = {pr[f"A{r}"].value: ev.value(f"N{r}", "Lab Proteins") for r in range(5, 10)}
    assert role["Binder A"] == "Ligand" and role["Control IgG"] == "Reference" and role["Analyte X"] == "Analyte"
    assert "'Plate Map'!" in pr["H5"].value


def test_a_build_never_overwrites_an_existing_output(demo_run):
    out = os.path.join(demo_run, "rebuild", "mine.prism")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    with open(out, "wb") as fh:
        fh.write(b"my graphs")
    code, stdout, err = run_cli("build", demo_run, "--no-template", "-o", out)
    assert code == 0, err
    assert open(out, "rb").read() == b"my graphs"
    assert os.path.exists(os.path.join(demo_run, "rebuild", "mine_2.prism")) and "already exists" in stdout


def test_build_from_scratch_gives_one_analysis_per_ligand(demo_run):
    import zipfile
    code, _, err = run_cli("build", demo_run, "--no-template")
    assert code == 0, err
    z = zipfile.ZipFile(os.path.join(demo_run, "rebuild", "Demo binding experiment.prism"))
    analyses = {n.split("/")[1] for n in z.namelist() if n.startswith("analyses/") and n.endswith("/results.json")}
    assert len(analyses) == 2


def test_kd_summary_follows_the_include_switches(demo_run):
    code, _, err = run_cli("workbook", demo_run)
    assert code == 0, err
    wb = openpyxl.load_workbook(os.path.join(demo_run, "rebuild", "Results.xlsx"))
    assert wb.sheetnames == ["Binder A - Equilibrium", "Binder B - Equilibrium", "Kinetic KDs", "KD summary"]
    ws, kd = wb["KD summary"], wb["Kinetic KDs"]
    assert ws["G6"].value == "='Kinetic KDs'!H7"
    before = X.Evaluator(wb).value("G6", "KD summary")
    det = next(r for r in range(1, kd.max_row + 1) if kd.cell(r, 1).value == "Fit rows") + 2
    for r in range(det, kd.max_row + 1):
        if kd.cell(r, 1).value == "Binder A" and kd.cell(r, 2).value == "Analyte X":
            kd.cell(r, 4).value = "No"
    assert before == pytest.approx(5, rel=0.2) and X.Evaluator(wb).value("G6", "KD summary") == "n/a"


def test_a_two_part_analyte_weighs_the_sum_of_its_parts():
    from bli2prism.workbook import analyte_mw
    an = [{"parts": ["P"], "mw": 20000}, {"parts": ["Q"], "mw": 30000}, {"parts": ["P", "Q"], "mw": 50000}]
    assert analyte_mw(an) == [20000.0, 30000.0, [0, 1]]       # the complex = the sum of the two single-protein analytes
