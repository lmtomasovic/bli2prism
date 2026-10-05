"""The demo experiment (simulated data) in examples/: reproducible, consistent with its generator, and a working run."""

import filecmp
import importlib.util
import os
import shutil

import openpyxl
import pytest

from bli2prism import workbook as W
from conftest import run_cli

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EXAMPLE = os.path.join(ROOT, "examples", "demo_binder_screen")


@pytest.fixture(scope="module")
def maker():
    spec = importlib.util.spec_from_file_location("make_example", os.path.join(ROOT, "examples", "make_example.py"))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


@pytest.fixture(scope="module")
def regenerated(maker, tmp_path_factory):
    d = str(tmp_path_factory.mktemp("example"))
    shutil.copy(os.path.join(EXAMPLE, "bli2prism_setup.xlsx"), d)          # the generator's input
    maker.make(d)
    return d


def test_the_committed_example_is_what_the_generator_writes(regenerated):
    names = sorted(f for f in os.listdir(EXAMPLE) if f.endswith((".xls", ".csv")))
    assert len(names) == 17
    for f in names:                                                   # deterministic: same seed, same bytes
        assert filecmp.cmp(os.path.join(EXAMPLE, f), os.path.join(regenerated, f), shallow=False), f
    assert openpyxl.load_workbook(os.path.join(EXAMPLE, "bli2prism_setup.xlsx")).sheetnames[0] == "bli2prism"


def test_the_example_runs_end_to_end_and_recovers_the_simulated_constants(regenerated, tmp_path):
    d = str(tmp_path / "run")
    shutil.copytree(EXAMPLE, d)
    for cmd in ("build", "workbook", "qc"):
        code, out, err = run_cli(cmd, d)
        assert code == 0, (cmd, out, err)
    ws = openpyxl.load_workbook(os.path.join(d, "rebuild", "Results.xlsx"), data_only=True)["KD summary"]
    kd = {(ws.cell(r, 1).value, ws.cell(r, 2).value): ws.cell(r, 7).value for r in range(6, 12)}
    truth = {("Binder B", "Analyte X"): 50, ("Binder B", "Analyte Y"): 20,
             ("Binder B", "Analyte Z"): 2000}
    for pair, nm in truth.items():
        assert kd[pair] == pytest.approx(nm, rel=0.15), pair           # mean per-row kinetic KD, nM
    assert kd[("Binder A", "Analyte Z")] == "n/a"                      # the non-binder has no usable kinetic row



def test_kd_summary_survives_a_fit_with_an_astronomical_interval():
    assert W._nM(400.0) == "n/a" and W._nM(float("nan")) == "n/a" and W._nM(-8.0) == pytest.approx(10.0)


def test_the_qc_report_names_only_the_folder_not_its_full_path(tmp_path):
    d = str(tmp_path / "run")
    shutil.copytree(EXAMPLE, d)
    out = str(tmp_path / "qc.html")
    code, _, err = run_cli("qc", d, "-o", out)
    assert code == 0, err
    html = open(out, encoding="utf-8").read()
    assert "Folder: <code>run</code>" in html and str(tmp_path) not in html
